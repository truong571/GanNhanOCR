"""Adapter ingest hai bản CHÉP TAY Vatican Borgiano Tonchinese → `prepared/_auto/<Sách>/` (TẬP ĐÁNH GIÁ chữ viết tay).

    SachKinhThayCaBinh = Borg.Tonch.18 (Philiphê Bỉnh, 529 trang chữ)
    SachDungLyHoThan   = Borg.Tonch.34 (112 trang chữ)

Vai trò giống hai bộ IHR-NomDB (ingest_ihr_book): sách CÓ NHÃN NGƯỜI (Excel phiên âm: `SinoNom_Char` +
`ChuQN_txt` theo câu) nên pipeline chạy lên nó CHỈ để ĐO — manifest đóng dấu `evaluation_only = true`.

⚠️ Adapter này **KHÔNG BAO GIỜ đọc** cột chữ Nôm người (`SinoNom_Char` của Excel, `prepared/<Sách>/nom_transcriptions/`,
trường `nom_*` của `prepared/<Sách>/transcriptions/*.json`). Từ Excel nó chỉ đọc 3 cột: `img_id`, `sentence_id`,
`ChuQN_txt` (vị trí cột kiểm theo tên tiêu đề, cột 3 bị bỏ qua khi đọc). Nhãn Nôm người chỉ dùng SAU KHI chạy, ở
`scripts/measure/borg_endtoend_eval.py`, để ĐO. Thư mục `prepared/<Sách>/` (ingest_borg_tonch, chứa nhãn người) không
bị đụng: đầu ra nằm ở `prepared/_auto/<Sách>/` (config: `paths.data_dir: prepared/_auto`).

Cách dựng (theo khuôn ingest_prose_book, hợp đồng engine y hệt các sách mới):
  1. Trang: nhóm dòng Excel theo `img_id` (sửa tên theo `typo_map` của ingest_borg_tonch), sắp theo số [NN] — ĐÚNG thứ tự
     và đánh số page_XXXX của ingest_borg_tonch (phép đo kiểm lại ánh xạ này). Ảnh gốc data/<Sách>/<img_id> (720 px, bản
     lớn nhất có trên máy) → pages/page_XXXX.png (xám, kéo tương phản `stretch`, CÙNG kích thước) + pages_denoised/.
  2. Kim OCR TOÀN TRANG (×1, gửi JPG gốc) với tham số books[].kim_* của config (lang_type 2 = Nôm), cache
     kim_raw/page_XXXX<hậu tố>.json theo (md5 ảnh gửi, tham số) — chạy lại = 0 lượt. Ngân sách `--budget`, sổ lượt gọi
     kim_calls.json, thử lại có giãn cách (`--retries`, ngoài phần thử lại HTTP của core.ocr.ocr_api), dừng ngay khi API
     báo lỗi TÀI KHOẢN (từ chối truy cập / đăng nhập hỏng / tài khoản không hoạt động / rơi về Guest) — không đoán tiếp.
     `--ocr cache` = chỉ dùng cache, KHÔNG BAO GIỜ gọi API (thiếu trang → dừng cứng, mã 3).
  3. Cột: hộp kim → chữ chia đều theo chiều cao hộp (expand_box_chars) → gom theo TÂM x (dung sai COL_TOL × bề ngang hộp
     trung vị) → gộp mảnh (≤ 2 chữ, sát cột kề) và đoạn cùng cột bị tách (chồng x ≥ 50 %, không chồng y); cột đọc PHẢI→TRÁI,
     trong cột trên→dưới.
  4. QN = âm tiết NGƯỜI của trang (các câu `ChuQN_txt` có img_id là trang ấy, làm sạch như ingest_borg_tonch, bỏ token
     không có chữ cái, chữ thường). Ghép chuỗi chữ kim cả trang với chuỗi âm bằng DP đơn điệu của ingest_prose_book
     (khớp +3 khi chữ ∈ R(âm) của Dict/QuocNgu_SinoNom.csv, 0 khi không tra được, chèn/xoá −1, hai đầu QN tự do); âm của
     từng cột = đoạn giữa âm đầu và âm cuối ghép với chữ của cột (âm kẽ chia trung điểm; âm thừa đầu/cuối trang gán cho
     cột đầu/cuối nếu ≤ max(2, 30 % số chữ cột)). Cột n_match < 3 hoặc tỉ lệ < 0,25 → dòng giữ chỗ `khongkhop` (REVIEW).
  5. Ghi đúng hợp đồng engine: detected/page_XXXX_ocr_cache.json (coords_space=fullpage, cột 1 = PHẢI nhất),
     transcriptions/page_XXXX.{txt,json} (1 dòng/cột, `dp_ratio` cho cổng (e) qn_count_gate), manifest.json (cổng + sổ API +
     evaluation_only). Trang không có kim (thất bại / ngoài ngân sách) KHÔNG được ghi (engine bỏ qua) và bị ghi cờ.

Chạy:
  .venv/bin/python -m pipeline.tools.ingest_borg_book --book SachDungLyHoThan --limit 1 --ocr kim --budget 1   # thử 1 lượt
  .venv/bin/python -m pipeline.tools.ingest_borg_book --book SachKinhThayCaBinh --ocr kim --budget 600
  .venv/bin/python -m pipeline.tools.ingest_borg_book --book SachKinhThayCaBinh --ocr cache      # 0 API (run_pipeline --book all)
  .venv/bin/python -m pipeline.tools.ingest_borg_book --status                                    # độ phủ cache 2 sách, 0 API
Test: .venv/bin/python -m pipeline.tools.ingest_borg_book --selftest  (không API, không ghi ngoài thư mục tạm)
"""
from __future__ import annotations

import argparse
import contextlib
import io
import json
import re
import statistics
import sys
import time
import unicodedata
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from pipeline.tools.ingest_lithograph_book import (  # noqa: E402  (lõi dùng chung)
    CONTENT_PLACEHOLDER, _nom_to_qn_readings, _rel, book_kim_params, box_rect, expand_box_chars, kim_cache_suffix,
    kim_params_of, prepare_image)
from pipeline.tools.ingest_prose_book import (  # noqa: E402  (DP đơn điệu chữ↔âm của văn xuôi)
    PROSE_GAP, PROSE_MATCH, PROSE_MIN_MATCH, PROSE_MIN_RATIO, PROSE_SUB, align_sequences, split_syllables_by_column)

# Chỉ dùng bảng tên tệp/sửa tên + hàm làm sạch QN của ingest_borg_tonch (không dùng hàm đọc Excel của nó: hàm ấy đọc cột Nôm).
from pipeline.tools.ingest_borg_tonch import BOOK_SPECS, clean_quocngu_syllables  # noqa: E402

BOOKS = {
    "SachKinhThayCaBinh": dict(ms="Borg.Tonch.18", n_pages_expect=529,
                               edition="Vatican, Borgiano Tonchinese 18 — Sách kinh Thầy cả Bỉnh (Philiphê Bỉnh), chép tay"),
    "SachDungLyHoThan": dict(ms="Borg.Tonch.34", n_pages_expect=112,
                             edition="Vatican, Borgiano Tonchinese 34 — Sách Dũng Lý Hộ Thần, chép tay"),
}
ALIASES = {"MSS_Borg_tonch_18": "SachKinhThayCaBinh", "MSS_Borg_tonch_34": "SachDungLyHoThan"}
XLSX_COLS = ("img_id", "sentence_id", "SinoNom_Char", "ChuQN_txt")   # tiêu đề Excel (kiểm, không đoán vị trí)
XLSX_READ = ("img_id", "sentence_id", "ChuQN_txt")                  # CHỈ ba cột này được đọc
COL_TOL = 0.6            # × bề ngang hộp kim trung vị: |tâm hộp − tâm cụm| tối đa để cùng cột
FRAG_MAX = 2             # cụm ≤ bấy nhiêu chữ = mảnh -> gộp vào cột kề nếu tâm cách < FRAG_DIST × bước cột
FRAG_DIST = 0.6
SEG_XOV = 0.5            # hai cụm chồng x ≥ 50 % bề ngang hẹp hơn và không chồng y -> một cột bị tách đôi
END_SLACK = (2, 0.3)     # âm thừa đầu/cuối trang gán cho cột đầu/cuối nếu ≤ max(2, 30 % số chữ cột)
COLS_EXPECT = (8, 13)    # cột/trang hợp lý (đã biết ~10–11): ngoài khoảng -> cờ
ACCOUNT_RE = re.compile(r"access_denied|Truy cập bị từ chối|not\s*activ|inactiv|Auto-login failed|No OCR token|"
                        r"expired|Guest Mode|tài khoản|chưa kích hoạt|bị khoá", re.I)
DEFAULT_OUT = REPO / "prepared" / "_auto"


class KimAccountError(RuntimeError):
    """API kim báo lỗi TÀI KHOẢN — dừng cả lượt (không gọi tiếp, không đoán)."""


class KimBudgetError(RuntimeError):
    """Hết ngân sách lượt gọi."""


def canon_book(book: str) -> str:
    return ALIASES.get(book, book)


# ---------------------------------------------------------------------------
# 1. Trang + QN người (KHÔNG đọc cột SinoNom_Char)
# ---------------------------------------------------------------------------
def _page_sort_key(img_id: str) -> tuple[int, str]:
    """== ingest_borg_tonch.parse_excel_groundtruth.page_sort_key (số trong [NN], rồi tên)."""
    m = re.search(r"\[(\d+)\]", img_id)
    return (int(m.group(1)), img_id) if m else (999999, img_id)


def read_qn_rows(xlsx: Path) -> list[dict]:
    """Excel → [{img_id, sentence_id, qn, row}] theo thứ tự dòng. Kiểm tiêu đề = XLSX_COLS; chỉ lấy ô của XLSX_READ."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr = tuple(str(h or "").strip() for h in next(it))
    if hdr[:len(XLSX_COLS)] != XLSX_COLS:
        raise SystemExit(f"[ingest_borg] {xlsx}: tiêu đề {hdr} ≠ {XLSX_COLS} — không đoán vị trí cột")
    idx = {name: hdr.index(name) for name in XLSX_READ}
    out = []
    for r, row in enumerate(it, start=2):
        row = tuple(row or ()) + (None,) * max(0, len(hdr) - len(row or ()))
        img = row[idx["img_id"]]
        if not img:
            continue
        out.append(dict(img_id=str(img).strip(), sentence_id=str(row[idx["sentence_id"]] or f"{img}.{r}"),
                        qn=str(row[idx["ChuQN_txt"]] or ""), row=r))
    wb.close()
    return out


def qn_syllables(text: str) -> list[str]:
    """Câu QN người → âm tiết: clean_quocngu_syllables (bỏ chú thích [a], tách gạch nối, chuẩn dấu) + bỏ token không có
    chữ cái + chữ thường NFC (khoá từ điển)."""
    out = []
    for t in clean_quocngu_syllables(text):
        if any(ch.isalpha() for ch in t):
            out.append(unicodedata.normalize("NFC", t.lower()))
    return out


def load_pages(book: str) -> list[dict]:
    """Trang theo thứ tự của ingest_borg_tonch → [{page, page_name, img_id, folio, file, sentences, syllables, rows}]."""
    spec = BOOK_SPECS[book]
    typo = spec.get("typo_map") or {}
    groups: dict[str, dict] = {}
    for r in read_qn_rows(Path(spec["data_dir"]) / spec["xlsx_name"]):
        img = typo.get(r["img_id"], r["img_id"])
        g = groups.setdefault(img, dict(img_id=img, sentences=[], rows=[]))
        g["sentences"].append(dict(sentence_id=r["sentence_id"], qn_raw=r["qn"], syllables=qn_syllables(r["qn"])))
        g["rows"].append(r["row"])
    recs = []
    for i, (img, g) in enumerate(sorted(groups.items(), key=lambda kv: _page_sort_key(kv[0])), start=1):
        fm = re.search(r"\]_([0-9a-zA-Z\.]+)\.jpg", img)
        syl = [s for sent in g["sentences"] for s in sent["syllables"]]
        recs.append(dict(page=i, page_name=f"page_{i:04d}", img_id=img, folio=fm.group(1) if fm else img,
                         file=Path(spec["data_dir"]) / img, sentences=g["sentences"], syllables=syl,
                         rows=(min(g["rows"]), max(g["rows"]))))
    return recs


# ---------------------------------------------------------------------------
# 2. Kim: cache + ngân sách + sổ + dừng khi lỗi tài khoản
# ---------------------------------------------------------------------------
class _Tee(io.TextIOBase):
    """Ghi qua stderr thật đồng thời giữ bản sao (để nhận lỗi tài khoản do core.ocr.ocr_api in ra)."""

    def __init__(self, real):
        self.real, self.buf = real, io.StringIO()

    def write(self, s):
        self.buf.write(s)
        return self.real.write(s)

    def flush(self):
        self.real.flush()


def load_kim_cache(raw_path: Path, image_hash: str, kim: dict) -> list[dict] | None:
    """Cache hợp lệ = trùng md5 ảnh gửi + tham số gọi (== ingest_lithograph_book.kim_boxes)."""
    if not raw_path.exists():
        return None
    try:
        d = json.loads(raw_path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return None
    if d.get("image_hash") == image_hash and isinstance(d.get("boxes"), list) and kim_params_of(d.get("kim_params")) == kim:
        return d["boxes"]
    return None


class KimClient:
    """Gọi kim qua core.ocr.ocr_api (upload_image + recognize); đếm lượt, ghi sổ, dừng khi lỗi tài khoản / hết ngân sách."""

    def __init__(self, ledger: Path, budget: int, retries: int = 2, backoff: float = 5.0, max_consec_fail: int = 3,
                 api=None, sleep=time.sleep):
        self.ledger, self.budget, self.retries, self.backoff = ledger, int(budget), int(retries), float(backoff)
        self.max_consec_fail, self.sleep = int(max_consec_fail), sleep
        self.api = api
        self.n_pages_sent = self.n_attempts = self.n_ok = self.n_fail = self.consec_fail = 0
        self.seconds = 0.0
        self.stopped = ""
        self._log: list[dict] = []

    def _ocr_api(self):
        if self.api is None:
            from core.ocr import ocr_api
            self.api = ocr_api
        return self.api

    def _one(self, src: Path, kim: dict) -> tuple[list[dict] | None, str]:
        api = self._ocr_api()
        tee = _Tee(sys.stderr)
        with contextlib.redirect_stderr(tee):
            fname = api.upload_image(str(src))
            boxes = api.recognize(fname, **kim) if fname else None
        err = tee.buf.getvalue()
        guest = bool(getattr(api, "is_guest_mode", lambda: False)())
        if guest or (boxes is None and ACCOUNT_RE.search(err)):
            return None, "account_error: " + (("guest_mode; " if guest else "") + " | ".join(
                ln.strip() for ln in err.splitlines() if ACCOUNT_RE.search(ln))[:400])
        if boxes is None:
            return None, "fail: " + (" | ".join(ln.strip() for ln in err.splitlines() if ln.strip())[-300:] or "?")
        return boxes, "ok"

    def call(self, page: str, src: Path, kim: dict) -> list[dict] | None:
        """Một TRANG = một lượt (upload + recognize). Thử lại `retries` lần khi lỗi thường; lỗi tài khoản -> KimAccountError."""
        if self.stopped:
            raise KimAccountError(self.stopped)
        if self.n_pages_sent >= self.budget:
            raise KimBudgetError(f"hết ngân sách {self.budget} lượt")
        self.n_pages_sent += 1
        for k in range(self.retries + 1):
            t0 = time.time()
            boxes, st = self._one(src, kim)
            dt = time.time() - t0
            self.seconds += dt
            self.n_attempts += 1
            self._log.append(dict(t=time.strftime("%Y-%m-%dT%H:%M:%S"), page=page, attempt=k + 1, status=st[:200],
                                  secs=round(dt, 2), n_boxes=None if boxes is None else len(boxes), kim=kim))
            self.flush()
            if st.startswith("account_error"):
                self.stopped = st
                raise KimAccountError(st)
            if boxes is not None:
                self.n_ok += 1
                self.consec_fail = 0
                return boxes
            if k < self.retries:
                self.sleep(self.backoff * (3 ** k))
        self.n_fail += 1
        self.consec_fail += 1
        if self.consec_fail >= self.max_consec_fail:
            self.stopped = f"{self.consec_fail} trang liên tiếp thất bại (không phải lỗi tài khoản) — dừng để người kiểm"
            raise KimAccountError(self.stopped)
        return None

    def flush(self):
        """Nối sổ lượt gọi (JSON lines) — không bao giờ ghi đè lịch sử."""
        if not self._log:
            return
        self.ledger.parent.mkdir(parents=True, exist_ok=True)
        with open(self.ledger, "a", encoding="utf-8") as f:
            for r in self._log:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        self._log = []

    def stats(self) -> dict:
        return dict(pages_sent=self.n_pages_sent, attempts=self.n_attempts, ok=self.n_ok, failed_pages=self.n_fail,
                    seconds=round(self.seconds, 1), budget=self.budget, stopped=self.stopped or None)


def ledger_totals(ledger: Path) -> dict:
    """Tổng sổ lượt gọi kim của sách (mọi lần chạy)."""
    n = ok = acc = 0
    secs = 0.0
    if ledger.exists():
        for ln in ledger.read_text(encoding="utf-8").splitlines():
            try:
                r = json.loads(ln)
            except Exception:  # noqa: BLE001
                continue
            n += 1
            ok += r.get("status") == "ok"
            acc += str(r.get("status", "")).startswith("account_error")
            secs += float(r.get("secs") or 0)
    return dict(attempts=n, ok=ok, account_errors=acc, seconds=round(secs, 1))


# ---------------------------------------------------------------------------
# 3. Hộp kim → cột (tâm x, phải→trái)
# ---------------------------------------------------------------------------
def cluster_columns(boxes: list[dict]) -> tuple[list[list[dict]], dict]:
    """Chữ kim → cột. Trả ([cột PHẢI→TRÁI: [chữ trên→dưới]], stats)."""
    items = []
    for b in boxes:
        chars = expand_box_chars(b)
        if not chars:
            continue
        x0, y0, x1, y1 = box_rect(b)
        items.append(dict(cx=(x0 + x1) / 2.0, x0=x0, x1=x1, y0=y0, y1=y1, w=max(1, x1 - x0), chars=chars))
    st = dict(n_boxes=len(boxes), n_boxes_text=len(items), n_chars=sum(len(i["chars"]) for i in items),
              n_frag_merged=0, n_seg_merged=0, box_w_med=None, col_pitch=None)
    if not items:
        return [], st
    multi = [i["w"] for i in items if len(i["chars"]) >= 2] or [i["w"] for i in items]
    wmed = float(statistics.median(multi))
    st["box_w_med"] = round(wmed, 1)
    tol = COL_TOL * wmed
    cl: list[dict] = []
    for it in sorted(items, key=lambda i: -i["cx"]):
        if cl and abs(it["cx"] - cl[-1]["cx"]) <= tol:
            c = cl[-1]
            n0, n1 = c["n"], len(it["chars"])
            c["cx"] = (c["cx"] * n0 + it["cx"] * n1) / (n0 + n1)
            c["n"] += n1; c["items"].append(it)
            c["x0"] = min(c["x0"], it["x0"]); c["x1"] = max(c["x1"], it["x1"])
            c["y0"] = min(c["y0"], it["y0"]); c["y1"] = max(c["y1"], it["y1"])
        else:
            cl.append(dict(cx=it["cx"], n=len(it["chars"]), items=[it], x0=it["x0"], x1=it["x1"], y0=it["y0"], y1=it["y1"]))

    def merge(a: dict, b: dict) -> dict:
        n = a["n"] + b["n"]
        return dict(cx=(a["cx"] * a["n"] + b["cx"] * b["n"]) / n, n=n, items=a["items"] + b["items"],
                    x0=min(a["x0"], b["x0"]), x1=max(a["x1"], b["x1"]), y0=min(a["y0"], b["y0"]), y1=max(a["y1"], b["y1"]))
    # đoạn cùng cột bị tách (tâm lệch > tol nhưng chồng x nhiều, không chồng y)
    changed = True
    while changed and len(cl) > 1:
        changed = False
        for k in range(len(cl) - 1):
            a, b = cl[k], cl[k + 1]
            ov = min(a["x1"], b["x1"]) - max(a["x0"], b["x0"])
            narrow = max(1, min(a["x1"] - a["x0"], b["x1"] - b["x0"]))
            yov = min(a["y1"], b["y1"]) - max(a["y0"], b["y0"])
            if ov >= SEG_XOV * narrow and yov <= 0:
                cl[k:k + 2] = [merge(a, b)]; st["n_seg_merged"] += 1; changed = True
                break
    cx = [c["cx"] for c in cl]
    pitch = float(statistics.median([cx[i] - cx[i + 1] for i in range(len(cx) - 1)])) if len(cx) >= 2 else float("inf")
    st["col_pitch"] = round(pitch, 1) if pitch != float("inf") else None
    # mảnh ≤ FRAG_MAX chữ sát cột kề
    changed = True
    while changed and len(cl) > 1:
        changed = False
        for k, c in enumerate(cl):
            if c["n"] > FRAG_MAX:
                continue
            nb = [j for j in (k - 1, k + 1) if 0 <= j < len(cl)]
            j = min(nb, key=lambda j: abs(cl[j]["cx"] - c["cx"]))
            if abs(cl[j]["cx"] - c["cx"]) < FRAG_DIST * pitch:
                lo, hi = sorted((k, j))
                cl[lo:hi + 1] = [merge(cl[lo], cl[hi])]; st["n_frag_merged"] += 1; changed = True
                break
    cols = []
    for c in cl:
        chars = [ch for it in c["items"] for ch in it["chars"]]
        chars.sort(key=lambda ch: ch["y_center"])
        cols.append(chars)
    return cols, st


# ---------------------------------------------------------------------------
# 4. DP chữ kim ↔ âm người của TRANG
# ---------------------------------------------------------------------------
def assign_page_syllables(cols_chars: list[list[str | None]], syls: list[str], n2q: dict[str, set[str]],
                          min_match: int = PROSE_MIN_MATCH, min_ratio: float = PROSE_MIN_RATIO) -> tuple[list[dict], dict]:
    """Cột (thứ tự đọc) → [{syllables|None, n_match, n_pair, n_chars, ratio, j0, j1, matched}], info."""
    flat = [ch for col in cols_chars for ch in col]
    pairs, score = align_sequences(flat, syls, n2q)
    spans = split_syllables_by_column([len(c) for c in cols_chars], pairs, len(syls))
    have = [k for k, s in enumerate(spans) if s["j0"] is not None]
    lead = trail = 0
    if have:
        k0, k1 = have[0], have[-1]
        lead = spans[k0]["j0"]
        if 0 < lead <= max(END_SLACK[0], END_SLACK[1] * len(cols_chars[k0])):
            spans[k0]["j0"] = 0
            lead = 0
        trail = len(syls) - spans[k1]["j1"]
        if 0 < trail <= max(END_SLACK[0], END_SLACK[1] * len(cols_chars[k1])):
            spans[k1]["j1"] = len(syls)
            trail = 0
    out, n_ok = [], 0
    for col, sp in zip(cols_chars, spans):
        n = len(col)
        ratio = sp["n_match"] / n if n else 0.0
        good = sp["j0"] is not None and sp["n_match"] >= min_match and ratio >= min_ratio and sp["j1"] > sp["j0"]
        n_ok += int(good)
        out.append(dict(syllables=(syls[sp["j0"]:sp["j1"]] if good else None), n_match=sp["n_match"], n_pair=sp["n_pair"],
                        n_chars=n, ratio=round(ratio, 3), j0=sp["j0"], j1=sp["j1"], matched=good))
    n_match_total = sum(p[2] for p in pairs)
    covered = sum((o["j1"] - o["j0"]) for o in out if o["matched"])
    info = dict(score=round(score, 1), n_chars=len(flat), n_syll=len(syls), n_pairs=len(pairs), n_match=n_match_total,
                match_ratio=round(n_match_total / len(flat), 3) if flat else None,
                syll_coverage=round(covered / len(syls), 3) if syls else None, lead_unassigned=lead,
                trail_unassigned=trail, n_cols=len(cols_chars), n_cols_matched=n_ok)
    return out, info


# ---------------------------------------------------------------------------
# 5. Chạy
# ---------------------------------------------------------------------------
def page_status(book: str, out_root: Path = DEFAULT_OUT, kim: dict | None = None, config: Path | None = None) -> dict:
    """Độ phủ cache kim của sách (0 API): {n_pages, n_cached, missing[:10], complete, kim, suffix}."""
    from core.ocr.ocr_api import _file_md5
    book = canon_book(book)
    kim = kim_params_of(kim or book_kim_params(book, config=config))
    sfx = kim_cache_suffix(kim)
    recs = [r for r in load_pages(book) if r["file"].exists()]
    miss = []
    for r in recs:
        raw = out_root / book / "kim_raw" / f"{r['page_name']}{sfx}.json"
        if load_kim_cache(raw, _file_md5(str(r["file"])), kim) is None:
            miss.append(r["page_name"])
    return dict(book=book, n_pages=len(recs), n_cached=len(recs) - len(miss), missing=miss[:10], n_missing=len(miss),
                complete=not miss and bool(recs), kim=kim, suffix=sfx, ledger=ledger_totals(out_root / book / "kim_calls.json"))


def ingest(book: str, pages: list[int] | None, limit: int | None, ocr: str, force: bool, out_root: Path,
           kim: dict | None, contrast: str = "stretch", budget: int = 0, retries: int = 2, backoff: float = 5.0,
           allow_missing: bool = False, client: KimClient | None = None, verbose: bool = True) -> dict:
    import cv2
    from core.image.image_processing import denoise_image
    from core.ocr.ocr_api import _file_md5, _pixel_hash, verify_cache_image

    book = canon_book(book)
    kim = kim_params_of(kim)
    kim_sfx = kim_cache_suffix(kim)
    recs = load_pages(book)
    if pages:
        recs = [r for r in recs if r["page"] in set(pages)]
    if limit:
        recs = recs[:limit]
    out = out_root / book
    dirs = {k: out / k for k in ("pages", "pages_denoised", "detected", "transcriptions", "kim_raw")}
    for d in dirs.values():
        d.mkdir(parents=True, exist_ok=True)
    ledger = out / "kim_calls.json"
    if client is None:
        client = KimClient(ledger, budget=budget, retries=retries, backoff=backoff)
    n2q = _nom_to_qn_readings()

    g = dict(n_pages_excel=len(recs), n_pages=0, n_pages_no_image=0, n_pages_no_kim=0, n_pages_cached=0,
             n_pages_new_kim=0, n_cols=0, n_cols_matched=0, n_cols_placeholder=0, n_cols_eq_n=0, n_chars_kim=0,
             n_chars_matched=0, n_syll_page=0, n_syll_assigned=0, n_boxes=0, n_frag_merged=0, n_seg_merged=0,
             n_cache_ok=0, cols_per_page={}, pages_cols_out_of_range=[], pages_no_kim=[], pages_no_image=[])
    results, flagged = [], {}
    stop_reason = ""
    missing_cache = []
    for rec in recs:
        name = rec["page_name"]
        flags: list[str] = []
        if not rec["file"].exists():
            g["n_pages_no_image"] += 1; g["pages_no_image"].append(rec["img_id"])
            flagged[name] = ["no_image"]
            continue
        src_hash = _file_md5(str(rec["file"]))
        raw_path = dirs["kim_raw"] / f"{name}{kim_sfx}.json"
        boxes = None if force else load_kim_cache(raw_path, src_hash, kim)
        if boxes is not None:
            g["n_pages_cached"] += 1
        elif ocr == "kim" and not stop_reason:
            try:
                boxes = client.call(name, rec["file"], kim)
            except KimAccountError as e:
                stop_reason = f"account: {e}"
            except KimBudgetError as e:
                stop_reason = f"budget: {e}"
            if boxes is not None:
                raw_path.write_text(json.dumps(dict(image=_rel(rec["file"]), image_hash=src_hash, kim_params=kim,
                                                    boxes=boxes), ensure_ascii=False, indent=1), encoding="utf-8")
                g["n_pages_new_kim"] += 1
        elif ocr == "cache":
            missing_cache.append(name)
        if boxes is None:
            g["n_pages_no_kim"] += 1; g["pages_no_kim"].append(name)
            flagged[name] = ["no_kim" + (f"({stop_reason[:60]})" if stop_reason else "")]
            continue

        png = dirs["pages"] / f"{name}.png"
        if force or not png.exists():
            prepare_image(rec["file"], png, contrast)
        den = dirs["pages_denoised"] / f"{name}.png"
        if force or not den.exists():
            cv2.imwrite(str(den), denoise_image(cv2.imread(str(png), cv2.IMREAD_GRAYSCALE)))
        img_hash = _file_md5(str(png))

        cols, cst = cluster_columns(boxes)
        g["n_boxes"] += cst["n_boxes"]; g["n_frag_merged"] += cst["n_frag_merged"]; g["n_seg_merged"] += cst["n_seg_merged"]
        if not (COLS_EXPECT[0] <= len(cols) <= COLS_EXPECT[1]):
            flags.append(f"n_cols={len(cols)}∉{list(COLS_EXPECT)}"); g["pages_cols_out_of_range"].append(name)
        res, info = assign_page_syllables([[c["char"] for c in col] for col in cols], rec["syllables"], n2q)
        columns_cache, cols_txt = [], []
        for k, (col, r) in enumerate(zip(cols, res), start=1):
            columns_cache.append([dict(char=c["char"], y_center=c["y_center"], bbox=c["bbox"]) for c in col])
            if r["matched"]:
                syl = list(r["syllables"]); raw = " ".join(syl)
                g["n_cols_matched"] += 1; g["n_chars_matched"] += r["n_match"]; g["n_syll_assigned"] += len(syl)
                g["n_cols_eq_n"] += int(len(syl) == len(col))
            else:
                syl = [CONTENT_PLACEHOLDER] * max(1, len(col)); raw = ""
                g["n_cols_placeholder"] += 1
                flags.append(f"col{k}:placeholder:match={r['n_match']}/{len(col)}")
            g["n_chars_kim"] += len(col)
            xs = [c["bbox"][0] for c in col] + [c["bbox"][2] for c in col]
            ys = [c["bbox"][1] for c in col] + [c["bbox"][3] for c in col]
            cols_txt.append(dict(column=k, reading_order=k, raw_text=raw, syllables=syl, num_syllables=len(syl),
                                 n_chars_kim=len(col), n_match=r["n_match"], dp_ratio=r["ratio"], syl_span=[r["j0"], r["j1"]],
                                 matched=bool(r["matched"]), x_range=[min(xs), max(xs)], y_range=[min(ys), max(ys)]))
        if info["lead_unassigned"] or info["trail_unassigned"]:
            flags.append(f"qn_ends_unassigned={info['lead_unassigned']}+{info['trail_unassigned']}")
        g["n_syll_page"] += len(rec["syllables"])
        cache = dict(image=_rel(png), image_hash=img_hash, pixel_hash=_pixel_hash(str(png)), framed=False, frame_pad=0,
                     coords_space="fullpage", n_columns=len(columns_cache), columns=columns_cache, boxes_raw=boxes,
                     layout="prose", box_source="kim", reading_direction="rtl", kim_source_image=_rel(rec["file"]),
                     kim_params=kim, col_pitch=cst["col_pitch"], source="borg_tonch", img_id=rec["img_id"])
        cache_path = dirs["detected"] / f"{name}_ocr_cache.json"
        cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
        vstat = verify_cache_image(str(cache_path), str(png))
        if vstat == "ok":
            g["n_cache_ok"] += 1
        else:
            flags.append(f"verify_cache={vstat}")
        (dirs["transcriptions"] / f"{name}.txt").write_text(
            "".join(" ".join(c["syllables"]) + "\n" for c in cols_txt), encoding="utf-8")
        (dirs["transcriptions"] / f"{name}.json").write_text(json.dumps(dict(
            book_page=rec["page"], page_name=name, source_file=rec["img_id"], folio=rec["folio"], columns=cols_txt,
            qn_line_confidences=[], qn_page_confidence=None, layout="prose", reading_direction="rtl",
            n_columns=len(cols_txt), qn_source="excel_ChuQN_txt", n_sentences=len(rec["sentences"]),
            sentence_ids=[s["sentence_id"] for s in rec["sentences"]], dp=info), ensure_ascii=False, indent=1),
            encoding="utf-8")
        g["n_pages"] += 1
        g["n_cols"] += len(cols_txt)
        g["cols_per_page"][name] = len(cols_txt)
        if flags:
            flagged[name] = flags
        results.append(dict(book_page=rec["page"], page_name=name, source_file=rec["img_id"], folio=rec["folio"],
                            excel_rows=list(rec["rows"]), num_columns=len(cols_txt), n_syll_page=len(rec["syllables"]),
                            total_syllables=sum(c["num_syllables"] for c in cols_txt),
                            ocr_chars=sum(len(c) for c in columns_cache), n_cols_matched=info["n_cols_matched"],
                            dp=info, kim_stats=cst, flags=flags))
        if verbose and (rec["page"] <= 3 or rec["page"] % 50 == 0):
            print(f"  {name} ({rec['img_id']}): {len(cols_txt)} cột, chữ kim {info['n_chars']}, âm {info['n_syll']}, "
                  f"khớp {info['n_match']} ({info['match_ratio']}), cột ghép {info['n_cols_matched']}/{len(cols_txt)}",
                  flush=True)

    if ocr == "cache" and missing_cache and not allow_missing:
        raise SystemExit(3)   # người gọi (main) đã in danh sách; không ghi manifest nửa vời
    cpp = list(g["cols_per_page"].values())
    g["cols_per_page_median"] = statistics.median(cpp) if cpp else None
    g["cols_per_page_hist"] = {str(k): cpp.count(k) for k in sorted(set(cpp))}
    g["match_ratio"] = round(g["n_chars_matched"] / g["n_chars_kim"], 3) if g["n_chars_kim"] else None
    g["cols_matched_ratio"] = round(g["n_cols_matched"] / g["n_cols"], 3) if g["n_cols"] else None
    g["cols_eq_n_ratio"] = round(g["n_cols_eq_n"] / g["n_cols_matched"], 3) if g["n_cols_matched"] else None
    g["syll_assigned_ratio"] = round(g["n_syll_assigned"] / g["n_syll_page"], 3) if g["n_syll_page"] else None
    n_img = g["n_pages_excel"] - g["n_pages_no_image"]
    g["kim_cache_complete"] = bool(n_img) and g["n_pages_no_kim"] == 0 and not (pages or limit)
    api = dict(this_run=client.stats(), ledger_total=ledger_totals(ledger), stop_reason=stop_reason or None)
    manifest = dict(book=book, source="vatican_borgiano_tonchinese", manuscript=BOOKS[book]["ms"],
                    edition=BOOKS[book]["edition"], layout="prose", n_columns="auto", reading_direction="rtl",
                    orig_dir=_rel(Path(BOOK_SPECS[book]["data_dir"])), scale=1, contrast=contrast, ocr=ocr, kim_src="orig",
                    kim_params=kim, kim_cache_suffix=kim_sfx, qn_source="excel ChuQN_txt (phiên âm người, theo trang)",
                    dp=dict(match=PROSE_MATCH, sub=PROSE_SUB, gap=PROSE_GAP, min_match=PROSE_MIN_MATCH,
                            min_ratio=PROSE_MIN_RATIO, end_slack=list(END_SLACK), col_tol=COL_TOL),
                    evaluation_only=True,
                    evaluation_note=("Bản CHÉP TAY có nhãn Nôm người (Excel SinoNom_Char). Nhãn pipeline sinh ra từ bộ này CHỈ "
                                     "để ĐO độ đúng trên chữ viết tay, KHÔNG trộn vào tập huấn luyện. Adapter chỉ đọc "
                                     "img_id/sentence_id/ChuQN_txt của Excel; không đọc cột SinoNom_Char, "
                                     "prepared/<Sách>/nom_transcriptions/ hay trường nom_* nào."),
                    gt_file=_rel(Path(BOOK_SPECS[book]["data_dir"]) / BOOK_SPECS[book]["xlsx_name"]),
                    subset=dict(pages=pages, limit=limit) if (pages or limit) else None,
                    pages=results, total_pages=len(results), total_syllables=sum(r["total_syllables"] for r in results),
                    api=api, gates=dict(**g, pages_flagged=flagged, n_pages_flagged=len(flagged)))
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    return manifest


def _parse_pages(s: str | None) -> list[int] | None:
    if not s:
        return None
    out: list[int] = []
    for tok in s.split(","):
        tok = tok.strip()
        if "-" in tok:
            a, b = tok.split("-", 1)
            out += list(range(int(a), int(b) + 1))
        elif tok:
            out.append(int(tok))
    return out or None


# ---------------------------------------------------------------------------
# 6. Selftest (không API, thư mục tạm)
# ---------------------------------------------------------------------------
def selftest() -> int:
    import tempfile
    n = ok = 0

    def chk(name, cond):
        nonlocal n, ok
        n += 1; ok += bool(cond)
        print(f"  {'PASS' if cond else 'FAIL'} {name}")

    def box(x0, y0, x1, y1, text):
        return dict(points=[[x0, y0], [x1, y0], [x1, y1], [x0, y1]], transcription=text)
    # --- gom cột: 3 cột phải→trái, một cột bị tách đôi dọc, một mảnh 1 chữ sát cột
    B = [box(600, 10, 650, 210, "甲乙丙丁"), box(500, 12, 548, 110, "戊己"), box(503, 120, 551, 220, "庚辛"),
         box(400, 10, 452, 260, "壬癸子丑寅"), box(455, 262, 470, 280, "卯")]
    cols, st = cluster_columns(B)
    chk("cột: 3 cột", len(cols) == 3)
    chk("cột: phải→trái", [c[0]["char"] for c in cols] == ["甲", "戊", "壬"])
    chk("cột: trên→dưới trong cột", "".join(c["char"] for c in cols[1]) == "戊己庚辛")
    chk("cột: mảnh 1 chữ gộp vào cột kề", "".join(c["char"] for c in cols[2]) == "壬癸子丑寅卯" and st["n_frag_merged"] == 1)
    chk("cột: rỗng", cluster_columns([])[0] == [])
    # --- DP: cột ghép âm, cột rác thành giữ chỗ
    n2q = {"甲": {"một"}, "乙": {"hai"}, "丙": {"ba"}, "丁": {"bốn"}, "戊": {"năm"}, "己": {"sáu"}, "庚": {"bảy"}, "辛": {"tám"}}
    syl = ["một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám"]
    res, info = assign_page_syllables([list("甲乙丙丁"), list("戊己庚辛"), list("ㄅㄆㄇ")], syl, n2q)
    chk("DP: cột 1 = 4 âm đầu", res[0]["syllables"] == syl[:4])
    chk("DP: cột 2 = 4 âm sau", res[1]["syllables"] == syl[4:])
    chk("DP: cột rác không ghép", res[2]["syllables"] is None and not res[2]["matched"])
    chk("DP: độ phủ âm = 1", info["syll_coverage"] == 1.0)
    res2, info2 = assign_page_syllables([list("乙丙丁"), list("戊己庚辛")], syl, n2q)
    chk("DP: âm thừa đầu trang (1) gán cột đầu", res2[0]["syllables"] == syl[:4] and info2["lead_unassigned"] == 0)
    # --- QN: làm sạch, bỏ số, chữ thường
    chk("QN: tách gạch nối + bỏ chú thích + chữ thường", qn_syllables("Phi-ri-tô [a] San-tô. 12 A-men.") ==
        ["phi", "ri", "tô", "san", "tô", "a", "men"])
    # --- Excel: chỉ đọc 3 cột; tiêu đề sai -> dừng
    try:
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.xlsx"
            wb = openpyxl.Workbook(); ws = wb.active
            ws.append(list(XLSX_COLS)); ws.append(["[2]_1v.jpg", "[2]_1v.1", "國語", "Quốc ngữ"])
            ws.append(["[1]_1r.jpg", "[1]_1r.1", "天主", "Chúa trời"]); wb.save(p)
            rows = read_qn_rows(p)
            chk("Excel: 2 dòng, không có khoá Nôm", len(rows) == 2 and all("SinoNom_Char" not in r and
                                                                              "國" not in json.dumps(r, ensure_ascii=False)
                                                                              for r in rows))
            wb = openpyxl.Workbook(); ws = wb.active; ws.append(["a", "b", "c", "d"]); wb.save(p)
            try:
                read_qn_rows(p); bad = False
            except SystemExit:
                bad = True
            chk("Excel: tiêu đề lạ -> dừng", bad)
    except ImportError:
        chk("openpyxl có sẵn", False)
    # --- KimClient: ngân sách, thử lại, lỗi tài khoản dừng, sổ nối thêm

    class FakeApi:
        def __init__(self, seq):
            self.seq, self.i, self.guest = list(seq), 0, False

        def upload_image(self, path):
            kind = self.seq[min(self.i, len(self.seq) - 1)]
            self.i += 1
            if kind == "deny":
                print("[OCR] Upload failed: {'title': 't_access_denied'}", file=sys.stderr); return None
            if kind == "fail":
                print("[OCR] Upload failed after 4 attempts: HTTP 503", file=sys.stderr); return None
            if kind == "guest":
                self.guest = True
            return "f.jpg"

        def recognize(self, fname, **k):
            return [box(0, 0, 10, 20, "甲乙")]

        def is_guest_mode(self):
            return self.guest
    with tempfile.TemporaryDirectory() as td:
        led = Path(td) / "calls.json"
        c = KimClient(led, budget=2, retries=1, backoff=0, api=FakeApi(["ok", "ok", "ok"]), sleep=lambda s: None)
        chk("kim: lượt 1 ok", c.call("p1", Path("x"), {}) is not None)
        c.call("p2", Path("x"), {})
        try:
            c.call("p3", Path("x"), {}); over = False
        except KimBudgetError:
            over = True
        chk("kim: vượt ngân sách -> dừng", over and c.n_pages_sent == 2)
        c2 = KimClient(led, budget=5, retries=2, backoff=0, api=FakeApi(["fail", "ok"]), sleep=lambda s: None)
        chk("kim: lỗi thường -> thử lại rồi ok", c2.call("p4", Path("x"), {}) is not None and c2.n_attempts == 2)
        c3 = KimClient(led, budget=5, retries=2, backoff=0, api=FakeApi(["deny"]), sleep=lambda s: None)
        try:
            c3.call("p5", Path("x"), {}); acc = False
        except KimAccountError:
            acc = True
        chk("kim: từ chối truy cập -> KimAccountError, không thử lại", acc and c3.n_attempts == 1)
        c4 = KimClient(led, budget=5, retries=0, backoff=0, api=FakeApi(["guest"]), sleep=lambda s: None)
        try:
            c4.call("p6", Path("x"), {}); g_ = False
        except KimAccountError:
            g_ = True
        chk("kim: rơi về Guest -> dừng", g_)
        c5 = KimClient(led, budget=9, retries=0, backoff=0, max_consec_fail=2, api=FakeApi(["fail"]), sleep=lambda s: None)
        c5.call("p7", Path("x"), {})
        try:
            c5.call("p8", Path("x"), {}); cf = False
        except KimAccountError:
            cf = True
        chk("kim: 2 trang liên tiếp hỏng -> dừng", cf)
        tot = ledger_totals(led)
        chk("sổ: nối thêm mọi lượt", tot["attempts"] == 8 and tot["account_errors"] == 2)
        # cache: đúng md5 + tham số
        rp = Path(td) / "page_0001_lt2.json"
        rp.write_text(json.dumps(dict(image_hash="h", kim_params=dict(ocr_id=1, lang_type=2, font_type=1),
                                      boxes=[box(0, 0, 1, 1, "甲")])), encoding="utf-8")
        k2 = kim_params_of(dict(lang_type=2))
        chk("cache: trùng md5 + tham số -> dùng", load_kim_cache(rp, "h", k2) is not None)
        chk("cache: khác md5 -> không dùng", load_kim_cache(rp, "x", k2) is None)
        chk("cache: khác tham số -> không dùng", load_kim_cache(rp, "h", kim_params_of(None)) is None)
    chk("hậu tố cache lt2", kim_cache_suffix(dict(lang_type=2)) == "_lt2")
    chk("bí danh MSS_Borg_tonch_18", canon_book("MSS_Borg_tonch_18") == "SachKinhThayCaBinh")
    print(f"ingest_borg_book selftest: {ok}/{n}")
    return 0 if ok == n else 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--book", choices=sorted(BOOKS) + sorted(ALIASES))
    ap.add_argument("--pages", default=None, help="vd 1,2,5-8 (số page_XXXX)")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--ocr", choices=["kim", "cache", "none"], default="cache",
                    help="kim = gọi API cho trang thiếu cache (có ngân sách) · cache = CHỈ cache, thiếu -> mã 3 · none = không kim")
    ap.add_argument("--budget", type=int, default=0, help="số TRANG tối đa được gửi kim trong lượt này (0 = không gọi)")
    ap.add_argument("--retries", type=int, default=2, help="thử lại mỗi trang khi lỗi thường (giãn 5 s, 15 s…)")
    ap.add_argument("--backoff", type=float, default=5.0)
    ap.add_argument("--allow-missing-cache", action="store_true", help="--ocr cache: bỏ trang thiếu cache thay vì dừng")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--out", default=str(DEFAULT_OUT), help="gốc đầu ra (mặc định prepared/_auto -> prepared/_auto/<Sách>)")
    ap.add_argument("--contrast", choices=["stretch", "otsu", "none"], default="stretch")
    ap.add_argument("--kim-config", default=None, help="config đọc books[].kim_* (mặc định config/pipeline_<Sách>.yaml)")
    ap.add_argument("--kim-lang-type", type=int, default=None, choices=[0, 1, 2])
    ap.add_argument("--kim-font-type", type=int, default=None, choices=[0, 1, 2])
    ap.add_argument("--status", action="store_true", help="in độ phủ cache kim (0 API) rồi thoát; mã 0 nếu đủ")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    out_root = Path(a.out)
    books = [canon_book(a.book)] if a.book else sorted(BOOKS)
    cfg = Path(a.kim_config) if a.kim_config else None
    if a.status:
        rc = 0
        for b in books:
            s = page_status(b, out_root, config=cfg)
            print(f"[ingest_borg] {b}: cache kim {s['n_cached']}/{s['n_pages']} trang ({s['suffix']}) · "
                  f"sổ: {s['ledger']['attempts']} lượt, ok {s['ledger']['ok']}, lỗi tài khoản {s['ledger']['account_errors']}"
                  + ("" if s["complete"] else f" · THIẾU {s['n_missing']} (vd {s['missing'][:3]})"))
            rc |= 0 if s["complete"] else 1
        return rc
    if not a.book:
        ap.error("--book bắt buộc (trừ --status/--selftest)")
    book = books[0]
    kim = book_kim_params(book, cli=dict(lang_type=a.kim_lang_type, font_type=a.kim_font_type), config=cfg)
    if a.ocr == "cache":
        s = page_status(book, out_root, kim=kim)
        if not s["complete"] and not a.allow_missing_cache and not (a.pages or a.limit):
            print(f"[ingest_borg] {book}: --ocr cache nhưng THIẾU cache kim {s['n_missing']}/{s['n_pages']} trang "
                  f"(vd {s['missing'][:3]}). KHÔNG gọi API. Chạy có xác nhận: ./run_pipeline.sh --book {book} "
                  f"(hoặc ingest_borg_book --ocr kim --budget N).", file=sys.stderr)
            return 3
    print(f"[ingest_borg] {book} · ocr={a.ocr} · budget={a.budget} · kim={kim} · cache kim_raw/*{kim_cache_suffix(kim)}.json "
          f"· out={out_root / book}", flush=True)
    try:
        m = ingest(book, _parse_pages(a.pages), a.limit, a.ocr, a.force, out_root, kim, a.contrast, a.budget, a.retries,
                   a.backoff, allow_missing=a.allow_missing_cache)
    except SystemExit as e:
        if e.code == 3:
            print(f"[ingest_borg] {book}: thiếu cache kim (--ocr cache) — dừng, không ghi manifest.", file=sys.stderr)
            return 3
        raise
    g, api = m["gates"], m["api"]
    print(f"[ingest_borg] {m['total_pages']} trang ghi / {g['n_pages_excel']} (không ảnh {g['n_pages_no_image']}, không kim "
          f"{g['n_pages_no_kim']}); cột {g['n_cols']} (trung vị/trang {g['cols_per_page_median']}), ghép {g['n_cols_matched']} "
          f"({g['cols_matched_ratio']}), M==N {g['cols_eq_n_ratio']}, chữ khớp {g['match_ratio']}, âm gán {g['syll_assigned_ratio']}; "
          f"kim: cache {g['n_pages_cached']}, mới {g['n_pages_new_kim']}, lượt này {api['this_run']}", flush=True)
    if api["stop_reason"]:
        print(f"[ingest_borg] DỪNG: {api['stop_reason']}", file=sys.stderr)
        return 2 if api["stop_reason"].startswith("account") else 4
    return 0


if __name__ == "__main__":
    sys.exit(main())
