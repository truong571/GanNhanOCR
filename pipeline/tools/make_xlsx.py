"""Xuất bộ nhãn ra .xlsx — bản đọc bằng Excel của chính `labels.csv`.

VÌ SAO KHÔNG DÙNG `pandas.to_excel`
-----------------------------------
Excel suy kiểu còn hăng hơn pandas, và kho này đã nhiều lần chảy máu vì đúng chuyện đó:
cột cờ '1' hoá 1.0 (label_in_train, thế hệ ≤25/08), âm Quốc ngữ thật "nan" hoá NaN,
`crop_w` '138' hoá 138.0, `image_md5` '0000e5…' mất số 0 đầu.
Nếu để Excel đoán thì `U+2029A` vẫn yên nhưng `[1453, 801, 1577, 942]` hay `0010` thì
không. Nên ở đây MẶC ĐỊNH là CHUỖI, và chỉ vài cột ĐO ĐƯỢC mới ghi kiểu số để còn lọc và
sắp xếp đúng. Cái giá là Excel hiện tam giác xanh "số lưu dạng chữ" ở vài ô — rẻ hơn
nhiều so với một cột nhãn hỏng âm thầm.

Từ 16/09 `labels.csv` giao nộp cố định 12 cột (A-9); công cụ này chạy trên bất kỳ CSV nào
(cả `labels_trace.csv`), nên COT_SO vẫn giữ các cột đo của sidecar.

Tệp .xlsx nằm ngoài git (`.gitignore: *.xlsx`) vì nó là ĐẦU RA dựng lại được. Công cụ này
tồn tại để nó luôn dựng lại được, thay vì là một tệp ai đó làm tay một lần rồi quên.

    python -m pipeline.tools.make_xlsx --labels re-dataset/labels.csv
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]

# Cột ĐO ĐƯỢC — ghi kiểu số để lọc/sắp xếp trong Excel cho đúng. Mọi cột khác giữ CHUỖI.
COT_SO = {"s3_cosine", "ink_pct", "stray_ink", "border_ink", "column"}

# Cột hay bị nhìn nhầm nên nới rộng sẵn (12 cột giao nộp + vài cột sidecar labels_trace.csv)
RONG = {"image": 34, "bbox": 22, "rule": 30, "image_md5": 15, "syllable": 11,
        "rule_goc": 22, "crop_quality_flag": 17, "context_evidence": 22, "box_source": 16}


def xuat(src: Path, dst: Path) -> dict:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with open(src, encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise SystemExit(f"{src} rỗng")
    cols = list(rows[0].keys())

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("labels")
    ws.append(cols)
    for r in rows:
        out = []
        for c in cols:
            v = r[c]
            if c in COT_SO and v not in ("", None):
                try:
                    out.append(float(v) if "." in v else int(v))
                except ValueError:
                    out.append(v)
            else:
                out.append(v)
        ws.append(out)
    wb.save(dst)

    # write_only không cho định dạng — mở lại để đông cứng tiêu đề + bật lọc
    wb = load_workbook(dst)
    ws = wb["labels"]
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E5266")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = RONG.get(c, 11)
    wb.save(dst)
    return {"dong": len(rows), "cot": len(cols), "mb": dst.stat().st_size / 1048576}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--labels", default="re-dataset/labels.csv")
    ap.add_argument("--out", default=None, help="mặc định: cùng chỗ, đuôi .xlsx")
    a = ap.parse_args(argv)
    src = Path(a.labels)
    if not src.exists():
        print(f"[xlsx] không có {src} — bỏ qua.")
        return 0
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("[xlsx] thiếu openpyxl — bỏ qua (pip install openpyxl).")
        return 0
    dst = Path(a.out) if a.out else src.with_suffix(".xlsx")
    r = xuat(src, dst)
    print(f"[xlsx] {dst} — {r['dong']:,} dòng × {r['cot']} cột · {r['mb']:.1f} MB")
    print(f"[xlsx] cột số: {', '.join(sorted(COT_SO))} — mọi cột khác giữ CHUỖI "
          f"(để Excel khỏi ăn 'nan', '0010', '1' -> 1.0)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
