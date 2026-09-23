"""Self-test cho pipeline/tools (không pytest).

Run:  .venv/bin/python -m pipeline.tools.selftest

Trước 2026-08-23 thư mục `tools/` KHÔNG có một test nào, dù `fix_tone` đã chạy thật và
`sem_score` được dùng để xếp hạng ô đáng chấm.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd

REPO = Path(__file__).resolve().parents[2]

_passed = 0
_failed = 0


def check(name: str, cond: bool, detail: str = "") -> None:
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  ok   {name}")
    else:
        _failed += 1
        print(f"  FAIL {name}  {detail}")


def _ds_out() -> Path:
    """Thư mục làm việc: dataset_out/ hoặc $DS_OUT (quy ước run_pipeline.sh N0a)."""
    import os
    return REPO / os.environ.get("DS_OUT", "dataset_out")


def _ds_dir() -> Path:
    """Thư mục GIAO NỘP: dataset/ (đã nạp phán quyết) rồi re-dataset/ (đem chấm).

    Ghim cứng "dataset" làm 18 test tự bỏ qua sau lần chạy đầu — đúng lỗi vừa vá ở
    update_bang_so_lieu. Bản thử nghiệm DS_OUT=dataset_out_v3 -> $DS_OUT/{dataset,re-dataset}
    (run_pipeline.sh:78-81), để kiểm schema 12 cột trước khi bộ v3 thay bộ thật.
    """
    root = _ds_out() if _ds_out() != REPO / "dataset_out" else REPO
    # 23/09: bộ STT nay ở dataset/SachThanhTruyen/ (bố cục đầu ra); giữ "dataset" cho cây cũ.
    for d in ("dataset/SachThanhTruyen", "dataset", "re-dataset"):
        if (root / d / "labels.csv").exists():
            return root / d
    return root / "re-dataset"


def _frame(pairs):
    return pd.DataFrame({"ocr_char": [c for c, _ in pairs],
                         "syllable": [s for _, s in pairs],
                         "tier": ["REVIEW"] * len(pairs)})


def test_fix_tone() -> None:
    print("[tools.fix_tone]")
    from pipeline.tools import fix_tone as F

    check("strip_tone bỏ THANH nhưng GIỮ dấu tạo chữ",
          F.strip_tone("trẩy") == "trây" and F.strip_tone("bụt") == "but",
          f'{F.strip_tone("trẩy")} / {F.strip_tone("bụt")}')
    check("strip_tone KHÔNG gộp vừa/vua (nếu gộp sẽ nối hai từ khác hẳn)",
          F.strip_tone("vừa") != "vua")
    check("strip_all bỏ MỌI dấu, kể cả đ -> d",
          F.strip_all("đến") == "den" and F.strip_all("ấy") == "ay",
          f'{F.strip_all("đến")} / {F.strip_all("ấy")}')

    # TẦNG 1 — chỉ lệch thanh, ứng viên duy nhất
    qn = {"trẩy": ["礼"], "bụt": ["孛"]}
    f, extra = F.plan(_frame([("礼", "trấy"), ("孛", "but")]), qn)
    check("tone_unique: sửa được cả hai", len(f) == 2, str(len(f)))
    check("tone_unique: đích đúng", sorted(f["syllable"]) == ["bụt", "trẩy"],
          str(sorted(f["syllable"])))
    check("nhãn luật là tone_unique", all(r.startswith("tone_unique") for r in f["rule"]),
          str(list(f["rule"])))

    # TẦNG 2 — rụng hẳn dấu tạo chữ (chỉ được thử khi tầng 1 không ra ứng viên)
    f2, _ = F.plan(_frame([("衣", "ay")]), {"ấy": ["衣"]})
    check("diacritic_unique: 'ay' -> 'ấy'",
          len(f2) == 1 and f2["syllable"].iloc[0] == "ấy", str(list(f2["syllable"])))
    f3, _ = F.plan(_frame([("旦", "den")]), {"đến": ["旦"]})
    check("diacritic_unique: 'den' -> 'đến' (qua đ/d)",
          len(f3) == 1 and f3["syllable"].iloc[0] == "đến", str(list(f3["syllable"])))

    # CHỐT AN TOÀN — âm là TỪ CÓ THẬT thì KHÔNG đụng (nhóm B: sai chữ hay sai thanh
    # là hai giả thuyết ngang nhau; sửa máy sẽ CHE mất lỗi chữ)
    f4, _ = F.plan(_frame([("異", "là")]), {"lạ": ["異"], "là": ["羅"]})
    check("KHÔNG đụng khi âm là từ có thật (nhóm B)", len(f4) == 0, str(len(f4)))

    # Nhập nhằng -> để nguyên, ghi vào báo cáo
    f5, e5 = F.plan(_frame([("係", "hê")]), {"hề": ["係"], "hệ": ["係"]})
    check("nhập nhằng thì KHÔNG sửa", len(f5) == 0)
    check("nhập nhằng được ghi vào báo cáo", e5["ambiguous"] and e5["ambiguous"][0]["n"] == 1,
          str(e5["ambiguous"]))

    # Corpus chốt: cùng chữ đã đi với MỘT ứng viên ở chỗ khác trong corpus
    df = _frame([("噲", "goi"), ("噲", "gọi"), ("噲", "gọi")])
    f6, _ = F.plan(df, {"gọi": ["噲"], "gỏi": ["噲"]})
    check("corpus chốt được khi 1 ứng viên đã có bằng chứng",
          len(f6) == 1 and f6["syllable"].iloc[0] == "gọi", str(list(f6["syllable"])))
    check("nhãn luật là *_corpus", all("_corpus" in r for r in f6["rule"]), str(list(f6["rule"])))

    # Rác marker: KHÔNG sửa, chỉ gắn cờ
    f7, e7 = F.plan(_frame([("白", "3"), ("景", "000"), ("祖", "%")]), {"ba": ["白"]})
    check("rác marker KHÔNG bị sửa", len(f7) == 0)
    check("rác marker được đếm", e7["marker_garbage_cells"] == 3,
          str(e7["marker_garbage_cells"]))

    # Luỹ đẳng: chạy trên bảng đã đúng thì không sửa gì
    f8, _ = F.plan(_frame([("礼", "trẩy")]), {"trẩy": ["礼"]})
    check("luỹ đẳng (âm đã khớp -> 0 sửa)", len(f8) == 0)


def test_sem_score() -> None:
    print("[tools.sem_score]")
    from pipeline.tools import sem_score as S

    # Corpus phải ĐỦ LỚN để IDF có nghĩa: từ xuất hiện ở 2/3 số mục thì
    # idf = log(3/3) = 0 và mọi điểm đều bằng 0 — đó là tính chất đúng của IDF,
    # không phải lỗi. Thêm nhiễu để "water" thành từ hiếm.
    uni = {"A": {"water", "river"}, "B": {"water", "stream"}, "C": {"fire", "flame"}}
    uni.update({f"N{i}": {f"filler{i}", "common"} for i in range(40)})
    qn = {"nuoc": ["A", "B"], "lua": ["C"]}
    sc = S.SemScorer(uni, qn)

    v_in = sc.score("A", "nuoc")     # leave-one-out: hồ sơ còn B (water, stream)
    v_out = sc.score("C", "nuoc")
    check("chữ cùng trường nghĩa -> điểm > 0", v_in and v_in > 0, str(v_in))
    check("chữ khác trường nghĩa -> điểm 0", v_out == 0, str(v_out))
    check("điểm cùng trường > khác trường", v_in > v_out)
    check("leave-one-out: KHÔNG tự khớp với chính mình",
          sc._profile("lua", exclude="C") == {} and sc.score("C", "lua") is None)
    check("chữ không có nghĩa Hán -> None", sc.score("ZZZ", "nuoc") is None)

    v, tag = sc.verdict("C", "nuoc")
    check("điểm 0 -> 'khong_ket_luan', KHÔNG phải 'sai'", tag == "khong_ket_luan", tag)
    check("ngưỡng xác nhận cao hơn ngưỡng khả năng", S.TAU_CONFIRM > S.TAU_LIKELY)


def test_syllable_normalize() -> None:
    """Chuẩn hoá TRƯỚC align — dùng đọc âm của MỌI chữ trong cột (cặp ghép chưa có)."""
    print("[align_engine.syllable_normalize]")
    from pipeline.align_engine import syllable_normalize as SN

    qn = {"trẩy": ["礼"], "ấy": ["衣"], "đến": ["旦"], "lạ": ["異"], "là": ["羅"],
          "hề": ["係"], "hệ": ["係"]}
    R = SN.build_readings(qn)
    keys = set(qn)
    ch = lambda *cs: [{"ocr_char": c} for c in cs]

    out, log = SN.normalize_column(ch("礼"), ["trấy"], keys, R)
    check("tone_unique: 'trấy' -> 'trẩy'", out == ["trẩy"], str(out))
    check("nhật ký ghi luật", log and log[0]["rule"] == "tone_unique", str(log))

    out, _ = SN.normalize_column(ch("衣"), ["ay"], keys, R)
    check("diacritic_unique: 'ay' -> 'ấy'", out == ["ấy"], str(out))
    out, _ = SN.normalize_column(ch("旦"), ["den"], keys, R)
    check("diacritic_unique qua đ/d: 'den' -> 'đến'", out == ["đến"], str(out))

    out, _ = SN.normalize_column(ch("異", "羅"), ["là"], keys, R)
    check("CHỐT 1 — âm là từ có thật thì KHÔNG đụng", out == ["là"], str(out))

    out, log = SN.normalize_column(ch("係"), ["hê"], keys, R)
    check("CHỐT 3 — nhập nhằng thì để nguyên", out == ["hê"], str(out))
    check("nhập nhằng được ghi log", any(e["action"] == "ambiguous_kept" for e in log), str(log))

    out, _ = SN.normalize_column(ch("礼"), ["3"], keys, R)
    check("rác marker không bị đụng (việc của parser)", out == ["3"], str(out))

    out, _ = SN.normalize_column(ch("礼"), ["xyz"], keys, R)
    check("không có ứng viên -> giữ nguyên", out == ["xyz"], str(out))

    out, _ = SN.normalize_column([], ["trấy"], keys, R)
    check("cột không có chữ -> giữ nguyên", out == ["trấy"], str(out))

    o1, _ = SN.normalize_column(ch("礼"), ["trấy"], keys, R)
    o2, l2 = SN.normalize_column(ch("礼"), o1, keys, R)
    check("luỹ đẳng (chạy lại không đổi thêm)",
          o2 == o1 and not [e for e in l2 if e["action"] == "fixed"], str(l2))
    check("build_readings có cache (cùng đối tượng -> cùng kết quả)",
          SN.build_readings(qn) is R)


def test_qn_charfix() -> None:
    """L3 — sửa lỗi ký tự OCR quốc ngữ trên âm NGOÀI từ điển (2026-09-24, mặc định TẮT)."""
    print("[align_engine.qn_charfix]")
    from pipeline.align_engine import qn_charfix as CF
    from pipeline.align_engine import syllable_normalize as SN

    K = {"cho", "chi", "con", "lạc", "khúc", "đã", "có", "éo", "eo", "cũng", "việc", "nó", "ô"}
    check("eho -> cho (e→c)", CF.fix_syllable("eho", K) == "cho")
    check("lạe -> lạc (e→c ở cuối)", CF.fix_syllable("lạe", K) == "lạc")
    check("ðã -> đã (ð→đ)", CF.fix_syllable("ðã", K) == "đã")
    check("‹ -> ∅ (xoá ký tự rác)", CF.fix_syllable("‹có", K) == "có")
    check("CHỐT 1 — âm ĐÃ trong từ điển thì KHÔNG đụng", CF.fix_syllable("cho", K) is None)
    check("CHỐT 2 — không có ứng viên -> None", CF.fix_syllable("zzz", K) is None)
    check("CHỐT 2 — nhiều ứng viên -> None (để nguyên)",
          CF.fix_syllable("ee", {"ce", "ec"}) is None, str(sorted(CF.char_variants("ee"))))
    check("bảng ký tự nạp được từ config/lexicon/qn_charfix.json",
          CF.CHAR_FIX.get("e") == "c" and CF.CHAR_FIX.get("ð") == "đ" and CF.CHAR_FIX.get("‹") == "",
          str(CF.CHAR_FIX))
    check("chỉ MỘT phép thay mỗi biến thể (không ghép)",
          "cho" in CF.char_variants("eho") and "chc" not in CF.char_variants("eho"),
          str(sorted(CF.char_variants("eho"))))

    qn = {"có": ["要"], "éo": ["要"], "cho": ["朱"]}
    R = SN.build_readings(qn)
    keys = set(qn)
    ch = lambda *cs: [{"ocr_char": c} for c in cs]
    out, log = SN.normalize_column(ch("要"), ["eó"], keys, R)
    check("MẶC ĐỊNH TẮT: charfix=False không đụng gì (hành vi cũ)", out == ["eó"], str(out))
    out, log = SN.normalize_column(ch("要"), ["eó"], keys, R, charfix=True)
    check("L3 CHẠY TRƯỚC tầng dời dấu: 'eó' -> 'có' (KHÔNG phải 'éo')", out == ["có"], str(out))
    check("nhật ký ghi kind=charfix", any(e.get("kind") == "charfix" for e in log), str(log))
    out, _ = SN.normalize_column(ch("朱"), ["cho"], keys, R, charfix=True)
    check("charfix không đụng âm hợp lệ", out == ["cho"], str(out))
    out, _ = SN.normalize_column(ch("朱"), ["3"], keys, R, charfix=True)
    check("charfix không đụng rác marker", out == ["3"], str(out))
    o1, _ = SN.normalize_column(ch("要"), ["eó"], keys, R, charfix=True)
    o2, l2 = SN.normalize_column(ch("要"), o1, keys, R, charfix=True)
    check("luỹ đẳng (chạy lại không sửa thêm)",
          o2 == o1 and not [e for e in l2 if e["action"] == "fixed"], str(l2))
    out, _ = SN.normalize_column([], ["eó"], keys, R, charfix=True)
    check("cột không có chữ: L3 vẫn chạy (quyết định ở phía QN, không cần kim)", out == ["có"], str(out))


def test_check_deps() -> None:
    """Kiểm thư viện chạy MỖI LẦN, và tôn trọng ba cạm bẫy của môi trường 3.14."""
    from pathlib import Path as _P
    REPO = _P(__file__).resolve().parents[2]
    print("[kiểm thư viện]")
    from pipeline.tools import check_deps as CD
    req = CD.parse_req()
    check(f"đọc được requirements.txt ({len(req)} gói)", len(req) >= 20)
    names = {n.lower() for n, _, _ in req}
    check("dùng pygame-ce, KHÔNG dùng pygame (3.14 không có wheel)",
          "pygame-ce" in names and "pygame" not in names)
    check("vietocr nằm trong nhóm phải cài --no-deps", "vietocr" in CD.NO_DEPS)
    a = CD.audit()
    check("mọi gói khai báo đều được phân loại",
          len(a["ok"]) + len(a["thieu"]) + len(a["lech"]) == len(req))
    src = (REPO / "pipeline" / "tools" / "check_deps.py").read_text(encoding="utf-8")
    check("--fix cài --no-deps RIÊNG chứ không gộp", '"--no-deps"' in src)
    check("--fix kiểm LẠI sau khi cài", "b = audit()" in src)
    rp = (REPO / "run_pipeline.sh").read_text(encoding="utf-8")
    check("run_pipeline gọi check_deps ở preflight", "pipeline.tools.check_deps" in rp)
    check("thiếu gói thì TỰ VÁ chứ không chỉ báo", "check_deps --fix" in rp)
    check("có đường thoát SKIP_DEPS", "SKIP_DEPS" in rp)


def test_dict_candidates() -> None:
    """Bảng ứng viên từ điển PHẢI có đủ cột phân xử.

    Bản đầu (2026-08-25) bỏ mất `similar_hit` và `char_in_dict` — đúng hai cột dùng để
    phân biệt "TỪ ĐIỂN THIẾU" với "OCR NHẦM TỰ DẠNG" — khiến người duyệt không có căn
    cứ nào để phán. Đây là test chặn việc đó tái diễn.
    """
    from pathlib import Path as _P
    import csv as _csv
    REPO = _P(__file__).resolve().parents[2]
    print("[ứng viên từ điển — đủ cột phân xử]")
    f = REPO / "docs" / "UNG_VIEN_MO_RONG_TU_DIEN.csv"
    if not f.exists():
        print("  [bỏ qua] chưa sinh bảng"); return
    head = next(_csv.reader(open(f, encoding="utf-8")))
    for c in ("similar_hit", "char_in_dict", "gold_count", "nghieng_ve"):
        check(f"có cột {c}", c in head)
    check("KHÔNG dùng cột ảnh mẫu (ô REVIEW không có crop)", "anh_mau_1" not in head)
    check("có toạ độ để mở ảnh trang gốc", "vi_tri_1" in head)
    rows = list(_csv.DictReader(open(f, encoding="utf-8")))
    check("có dòng", len(rows) > 0)
    # bất biến: hễ similar_hit KHÔNG rỗng thì phải xếp vào nhóm nghi OCR nhầm
    sai = [r for r in rows if r["similar_hit"].strip() and r["nghieng_ve"] != "B_OCR_NHAM"]
    check("similar_hit không rỗng => luôn xếp B_OCR_NHAM", not sai,
          f"{len(sai)} dòng lệch")
    check("cột người duyệt để TRỐNG", all(not r["NGUOI_DUYET"].strip() for r in rows))


def test_variant_table() -> None:
    """Bảng dị thể chỉ ĐỀ XUẤT, KHÔNG được sửa bộ nhãn."""
    from pathlib import Path as _P
    import csv as _csv
    REPO = _P(__file__).resolve().parents[2]
    print("[dị thể — máy đề xuất, người quyết]")
    src = (REPO / "pipeline" / "tools" / "variant_table.py").read_text(encoding="utf-8")
    check("KHÔNG ghi vào bộ nhãn", "labels.csv\", \"w\"" not in src and "to_csv" not in src)
    check("đòi HAI điều kiện: cùng âm VÀ nhìn giống", "oneway" in src and "mutual" in src)
    check("bỏ cặp KHÔNG nhìn giống (hai chữ khác nhau thật)", "if not oneway:" in src)
    f = REPO / "docs" / "UNG_VIEN_CHUAN_HOA_DI_THE.csv"
    if not f.exists():
        print("  [bỏ qua] chưa sinh bảng"); return
    rows = list(_csv.DictReader(open(f, encoding="utf-8")))
    check("mã trội luôn nhiều ô hơn mã phụ",
          all(int(r["so_o_troi"]) >= int(r["so_o_phu"]) for r in rows))
    check("cột người duyệt để TRỐNG", all(not r["NGUOI_DUYET"].strip() for r in rows))


def test_dataset_docs() -> None:
    """Tài liệu bộ giao nộp: số ĐỌC TỪ nhãn, KHÔNG bịa lai lịch thư tịch, và schema 12 cột.

    Từ 16/09 (A-9/A-10 giai đoạn 1) `labels.csv` giao nộp cố định 12 cột, không split,
    kèm sidecar labels_trace.csv + columns.csv. Bộ thế hệ cũ (≤25/08, 30+ cột, đang đóng
    băng ở re-dataset/) chỉ chạy phần kiểm không phụ thuộc schema — chạy
    `DS_OUT=dataset_out_v3 python -m pipeline.tools.selftest` để kiểm bản mới.
    """
    from pathlib import Path as _P
    REPO = _P(__file__).resolve().parents[2]
    print("[tài liệu bộ giao nộp]")
    ds = _ds_dir()
    if not (ds / "labels.csv").exists():
        print("  [bỏ qua] chưa có bộ giao nộp"); return
    for n in ("README.md", "DATASHEET.md", "NGUON_THU_TICH.md", "LICENSE.md"):
        check(f"có {n}", (ds / n).exists())
    rd = (ds / "README.md").read_text(encoding="utf-8")
    import csv as _csv
    rows = list(_csv.DictReader(open(ds / "labels.csv", encoding="utf-8")))
    hdr = list(rows[0].keys()) if rows else []
    nchar = sum(1 for r in rows if (r.get("label") or "").strip())
    check(f"README nêu đúng số nhãn ký tự ({nchar:,})", f"{nchar:,}" in rd)
    check("README CẢNH BÁO đừng gộp hai loại nhãn", "Đừng phát biểu" in rd)
    ng = (ds / "NGUON_THU_TICH.md").read_text(encoding="utf-8")
    check("lai lịch để TRỐNG chứ không bịa", "⬜ CHƯA ĐIỀN" in ng)

    # ---- SCHEMA 12 CỘT (16/09) ----------------------------------------------------
    # HỒI QUY cũ: README từng khẳng định "chia tách theo TRANG, không trang nào ở hai phía"
    # trong khi ĐO ĐƯỢC 360/444 trang có cột ở nhiều phía. Nay bộ giao nộp KHÔNG chia
    # train/val/test nữa (ràng buộc 16/09) nên README phải nói đúng điều đó và ghi công
    # thức hash cũ cho ai cần tự chia theo trang.
    from pipeline.export_final_dataset import GIAO_NOP, TRACE
    the_he_moi = hdr == GIAO_NOP
    if not the_he_moi:
        print(f"  [thế hệ cũ] {ds.relative_to(REPO)}/labels.csv có {len(hdr)} cột (bộ ≤25/08 "
              f"đang đóng băng) — bỏ qua phép kiểm schema 12 cột; kiểm bản mới bằng "
              f"DS_OUT=dataset_out_v3")
    else:
        check("labels.csv đúng 12 cột, đúng thứ tự GIAO_NOP", hdr == GIAO_NOP, str(hdr))
        _bo = {"split", "split_group", "label_in_train", "label_level", "usable_image",
               "page_cot_lech", "seg_backend", "readmitted_from_s3_demotion"}
        check("không còn cột đã bỏ (split/label_level/usable_image/page_cot_lech…)",
              not (_bo & set(hdr)), str(_bo & set(hdr)))
        check("image là khoá chính (không trùng)", len({r["image"] for r in rows}) == len(rows))
        check("tier chỉ GOLD/SILVER/SYLLABLE", {r["tier"] for r in rows} <= {"GOLD", "SILVER", "SYLLABLE"},
              str({r["tier"] for r in rows}))
        check("SYLLABLE <=> label rỗng (label_level suy được từ tier)",
              all((r["tier"] == "SYLLABLE") == (not r["label"].strip()) for r in rows))
        # sidecar labels_trace.csv: cùng số dòng, cùng thứ tự image, chỉ cột trong TRACE
        tp = ds / "labels_trace.csv"
        check("có labels_trace.csv", tp.exists())
        if tp.exists():
            tr = list(_csv.DictReader(open(tp, encoding="utf-8")))
            th = list(tr[0].keys()) if tr else []
            check("trace cùng số dòng và cùng thứ tự `image`",
                  [r["image"] for r in tr] == [r["image"] for r in rows],
                  f"{len(tr):,} vs {len(rows):,}")
            check("trace chỉ gồm cột trong TRACE, đúng thứ tự",
                  th == [c for c in TRACE if c in th], str(th))
            check("trace có crop_quality_flag (thay usable_image)", "crop_quality_flag" in th)
            for c in ("l1_tie", "qd01_locked", "qd01_excluded", "flank_gold"):
                if c in th:
                    check(f"trace cột cờ {c} ghi DÀY (không rỗng, không '1.0')",
                          {r[c] for r in tr} <= {"0", "1", "2"}, str(sorted({r[c] for r in tr}))[:40])
        # columns.csv: khoá book,page,column, phủ mọi cột của labels.csv
        cp = ds / "columns.csv"
        check("có columns.csv", cp.exists())
        if cp.exists():
            co = list(_csv.DictReader(open(cp, encoding="utf-8")))
            ch = list(co[0].keys()) if co else []
            check("columns.csv bắt đầu bằng khoá book,page,column", ch[:3] == ["book", "page", "column"], str(ch))
            _k = {(r["book"], r["page"], r["column"]) for r in co}
            check("columns.csv không trùng khoá", len(_k) == len(co))
            _lab = {(r["book"], r["page"], r["column"]) for r in rows}
            check(f"columns.csv phủ mọi cột của labels.csv ({len(_lab):,})", _lab <= _k,
                  f"thiếu {len(_lab - _k)}")
        check("README nói KHÔNG chia train/val/test", "Không chia train/val/test" in rd)
        check("README ghi công thức hash cũ để ai cần tự chia theo trang",
              "% 100" in rd and "md5" in rd and "< 80" in rd and "< 90" in rd)
        check("README KHÔNG còn hứa 'rời nhau theo TRANG' / cột split",
              "rời nhau theo TRANG" not in rd and "`split` / `split_group`" not in rd)
        check("README mô tả đủ 12 cột", all(f"| `{c}` |" in rd for c in GIAO_NOP))
        check("README nêu labels_trace.csv và columns.csv",
              "labels_trace.csv" in rd and "columns.csv" in rd)
        _syl = sum(1 for r in rows if not (r.get("label") or "").strip())
        if _syl:
            check("README cảnh báo lọc lớp-có-trong-train trong phạm vi GOLD (khỏi vứt SYLLABLE)",
                  'tier == "GOLD"' in rd and "vứt" in rd)

    # ---- HỒI QUY 2026-08-25: DATASHEET nói NGƯỢC README, và chốt chặn không hề đọc nó ----
    # Test cũ chỉ soi biến `rd` (=README) rồi báo 10/10 xanh, trong khi DATASHEET cùng thư
    # mục, sinh cùng lần chạy, khẳng định "chia tách neo ở mức CỘT — 360/444 trang có cột ở
    # nhiều phía". Câu đó là CHUỖI GHIM CỨNG sót lại từ trước khi đổi sang chia theo trang,
    # nằm lọt giữa một f-string mà mọi số quanh nó đều động, và nằm đúng mục "🔴 Giới hạn —
    # đọc trước khi dùng" mà run_pipeline.sh chỉ người đọc tới. Bộ giao nộp tự mâu thuẫn.
    dsh = (ds / "DATASHEET.md").read_text(encoding="utf-8")
    check("DATASHEET không còn chuỗi ghim cứng '360/444'", "360/444" not in dsh)
    if the_he_moi:
        check("DATASHEET không mâu thuẫn README: nói KHÔNG chia train/val/test",
              "Không chia train/val/test" in dsh and "label_in_train" not in dsh)
        check("DATASHEET không còn nhắc cột đã bỏ (usable_image/page_cot_lech)",
              "usable_image" not in dsh and "page_cot_lech" not in dsh)
    import re as _re
    _so_muc = _re.findall(r"^(\d+)\. \*\*", dsh, _re.M)
    check("DATASHEET không đánh trùng số mục", len(_so_muc) == len(set(_so_muc)),
          f"trùng: {[x for x in set(_so_muc) if _so_muc.count(x) > 1]}")
    _lab = [r["label"] for r in rows if (r.get("label") or "").strip()]
    _ngoai = sum(1 for c in _lab if not ("\u4e00" <= c <= "\u9fff"))
    _pct = f"{100*_ngoai/len(_lab):.2f}".replace(".", ",")
    check(f"DATASHEET nêu tỷ lệ ngoài CJK ĐO ĐƯỢC ({_pct}%)", f"{_pct}%" in dsh,
          "còn ghim cứng 1,63%?" if "1,63%" in dsh else "không thấy số đo")

    check("nói rõ vì sao để trống", "cố ý để trống" in ng.lower() or "bịa" in ng)


def test_proto_index_ro_ri_split() -> None:
    """Nguyên mẫu thị giác S3 không được lấy từ trang val/test.

    HỒI QUY 2026-08-25: chia tách chuyển từ mức CỘT sang mức TRANG, `index.csv` không ai
    sinh lại, và 8.615/41.835 nguyên mẫu (20,6%) gắn `split=train` hoá ra nằm trên trang
    NAY thuộc val/test. Phép kiểm `--check` cũ chỉ so TIỀN TỐ SÁCH nên vẫn báo "cùng thế
    hệ" — bảo đảm của build_rows bị phá mà không một tín hiệu nào. Trấn an sai, đúng lớp
    lỗi tệ nhất trong kho này.
    """
    import csv as _csv, tempfile
    from pipeline.tools import rebuild_proto_index as rpi
    print("[nguyên mẫu S3 vs chia tách]")
    lech, tong = rpi.split_lech()
    check("index.csv hiện hành: 0 nguyên mẫu rơi vào val/test", lech == 0,
          f"{lech:,}/{tong:,}")
    check("index.csv không rỗng", tong > 0, str(tong))

    # phép kiểm phải THẬT SỰ bắt được, không chỉ đọc số 0 rồi gật
    with tempfile.TemporaryDirectory() as td:
        d = Path(td)
        (d / "idx.csv").write_text(
            "path,label,unicode,split,source\n"
            "dataset_out/gold/stt2_page_0001_c01_001.png,A,U+0041,train,crop\n"
            "dataset_out/gold/stt2_page_0002_c01_001.png,B,U+0042,train,crop\n",
            encoding="utf-8")
        (d / "lab.csv").write_text(
            "book,page,split\nstt2,page_0001,train\nstt2,page_0002,test\n", encoding="utf-8")
        l2, t2 = rpi.split_lech(d / "lab.csv", d / "idx.csv")
        check("tiêm 1 nguyên mẫu val/test -> phép kiểm BẮT được", (l2, t2) == (1, 2),
              f"{l2}/{t2}")


def test_publish_doc_csv_phong_ve() -> None:
    """publish/ đọc CSV phải phòng vệ y như remediation/ — 'nan' là âm Việt thật."""
    print("[publish đọc CSV]")
    src = (REPO / "pipeline" / "publish" / "cli.py").read_text(encoding="utf-8")
    check("publish/cli.py dùng keep_default_na=False", "keep_default_na=False" in src)
    check("publish/cli.py đọc cột cờ là chuỗi", '"label_in_train": str' in src)


def test_step1_khong_mat_trang() -> None:
    """Số trang IN trên bản quét không duy nhất — trùng thì KHÔNG được nuốt trang.

    HỒI QUY 2026-08-25: `page_name = f"page_{n:04d}"` với n do OCR đọc từ ảnh. Trùng thì
    trang sau giữ nguyên tên trang trước, rồi MỌI bước sau (`if not img_path.exists()`,
    cache OCR Nôm, ảnh QN tạm, cache OCR QN) đều thấy "đã có" nên bỏ qua — nó chép lại y
    nguyên trang trước còn trang thật thì biến mất. results vẫn +1 nên total_pages đếm
    LƯỢT chứ không đếm TRANG. Đo được 3 trang mất: STT2 page_0142 (trang PDF 132),
    STT11 page_0010 (trang PDF 28 và 220).
    """
    print("[bước 1 không nuốt trang]")
    src = (REPO / "pipeline" / "step1_extract.py").read_text(encoding="utf-8")
    check("có sổ tên trang đã dùng", "seen_names" in src)
    check("trùng thì TÁCH TÊN chứ không đè", "_p{page_idx:04d}" in src)
    check("và NÓI RA chứ không nuốt", "trùng SỐ TRANG IN" in src)
    check("Total báo cả số tên trang duy nhất", "tên trang duy nhất" in src)

    # luật đặt tên phải tất định và không bao giờ đụng nhau
    seen: dict[str, int] = {}
    ten = []
    for idx, bp in [(0, 10), (28, 10), (130, 142), (132, 142), (220, 10)]:
        nm = f"page_{bp:04d}"
        if nm in seen:
            nm = f"{nm}_p{idx:04d}"
        seen[nm] = idx
        ten.append(nm)
    check("5 lượt trùng số -> 5 tên PHÂN BIỆT", len(set(ten)) == 5, str(ten))
    check("trang đầu giữ nguyên tên gốc", ten[0] == "page_0010")
    check("trang sau mang hậu tố truy được về trang PDF", ten[1] == "page_0010_p0028")


def test_co_trang_lech_cot() -> None:
    """Trang không đủ 9 cột phải được NÊU RA, và phải đếm trên TOÀN BỘ hàng.

    Bố cục ván khắc luôn 9 cột. Lần chạy 2026-08-25 lộ ra 3 trang thiếu cột, trong đó
    stt4/page_0110 có OCR Nôm ĐỦ 9 cột nhưng chỉ 8 cột sống tới bước ghép — nên phép kiểm
    "9 cột" chạy sau extract (đếm cột OCR Nôm) cho nó đi qua. Bước xuất là chỗ duy nhất
    nhìn được cả hai phía.

    Từ 16/09 (schema 12 cột) cờ `page_cot_lech` KHÔNG còn trong labels.csv: danh sách trang
    suy từ `columns.csv` (dựng trên all_rows) và phải khớp CHÍNH XÁC với labels_final.csv.
    Thế hệ cũ (≤25/08) vẫn kiểm cột `page_cot_lech` như trước.

    BẪY: nếu đếm cột trên hàng ĐÃ LỌC (bộ giao nộp bỏ REVIEW/SILVER) thì một trang lành
    có nguyên một cột rơi vào REVIEW sẽ bị đếm hụt và mang tiếng oan. Phải đếm trên all_rows.
    """
    import csv as _csv, collections as _cl
    print("[cờ trang lệch cột]")
    ds = _ds_dir()
    full = _ds_out() / "labels_final.csv"
    if not (ds / "labels.csv").exists() or not full.exists():
        print("  [bỏ qua] chưa có bộ giao nộp"); return
    rows = list(_csv.DictReader(open(ds / "labels.csv", encoding="utf-8")))
    cot_full: dict = {}
    for r in _csv.DictReader(open(full, encoding="utf-8")):
        cot_full.setdefault((r["book"], r["page"]), set()).add(r["column"])
    that_lech = {k for k, v in cot_full.items() if len(v) != 9}
    trang_pub = {(r["book"], r["page"]) for r in rows}

    if "page_cot_lech" in (rows[0] if rows else {}):
        # thế hệ cũ
        check("giá trị chỉ '0'/'1'", {r["page_cot_lech"] for r in rows} <= {"0", "1"})
        co_co = {(r["book"], r["page"]) for r in rows if r["page_cot_lech"] == "1"}
        check("cờ khớp CHÍNH XÁC tập trang thiếu cột (đếm trên toàn bộ hàng)",
              co_co == (that_lech & trang_pub), f"cờ {sorted(co_co)} vs thật {sorted(that_lech)}")
    else:
        cp = ds / "columns.csv"
        check("schema 12 cột: có columns.csv để suy trang thiếu cột", cp.exists())
        if not cp.exists():
            return
        pg = _cl.defaultdict(set)
        for r in _csv.DictReader(open(cp, encoding="utf-8")):
            pg[(r["book"], r["page"])].add(r["column"])
        tu_cot = {k for k, v in pg.items() if len(v) != 9} & trang_pub
        check("columns.csv dựng trên TOÀN BỘ hàng (khớp labels_final.csv từng cột)",
              {k: v for k, v in pg.items()} == cot_full,
              f"{len(pg)} vs {len(cot_full)} trang")
        check("trang thiếu cột suy từ columns.csv khớp CHÍNH XÁC labels_final.csv",
              tu_cot == (that_lech & trang_pub), f"columns {sorted(tu_cot)} vs thật {sorted(that_lech)}")
        for b_, pg_ in sorted(tu_cot):
            for f, ten in ((ds / "README.md", "README"), (ds / "DATASHEET.md", "DATASHEET")):
                check(f"{ten} nêu đích danh trang thiếu cột {b_}/{pg_}",
                      f"{b_}/{pg_}" in f.read_text(encoding="utf-8"))

    # BẪY đếm trên hàng đã lọc: dựng một trang lành có 9 cột nhưng chỉ 8 cột lọt bộ xuất
    all_rows = [{"book": "b", "page": "p", "column": str(i), "tier": "GOLD"} for i in range(1, 9)]
    all_rows.append({"book": "b", "page": "p", "column": "9", "tier": "REVIEW"})
    loc = [r for r in all_rows if r["tier"] == "GOLD"]
    dung = len({r["column"] for r in all_rows}) != 9
    sai = len({r["column"] for r in loc}) != 9
    check("đếm trên all_rows -> KHÔNG gắn cờ oan", dung is False)
    check("đếm trên hàng đã lọc -> gắn cờ OAN (bẫy đã tránh)", sai is True)

    for f, ten in ((ds / "README.md", "README"), (ds / "DATASHEET.md", "DATASHEET")):
        check(f"{ten} có nêu trang không đủ 9 cột",
              "không đủ 9 cột" in f.read_text(encoding="utf-8"))


def test_run_pipeline_grep_dem() -> None:
    """`grep -c` không khớp gì vẫn IN "0" rồi thoát mã 1 -> `|| echo 0` in thêm một "0".

    Biến thành "0\n0", và (( )) sặc: chính lỗi in ra ở preflight lần chạy 2026-08-25.
    """
    print("[preflight đếm bằng grep -c]")
    src = (REPO / "run_pipeline.sh").read_text(encoding="utf-8")
    check("không còn `grep -c ... || echo 0`",
          "|| echo 0)" not in src.replace("|| echo 0 )", "|| echo 0)"))
    check("dùng gán rồi mới chữa mã thoát", "|| n_old=0" in src and "|| n_new=0" in src)
    import subprocess
    r = subprocess.run(["bash", "-n", str(REPO / "run_pipeline.sh")], capture_output=True)
    check("run_pipeline.sh qua `bash -n`", r.returncode == 0, r.stderr.decode()[:200])
    # tái hiện lỗi gốc để chứng minh test không rỗng
    bad = subprocess.run(
        ["bash", "-c", 'n=$(grep -c zzz /dev/null || echo 0); [ "$n" = "0" ]'],
        capture_output=True)
    check("tái hiện được: `grep -c || echo 0` KHÔNG cho '0'", bad.returncode != 0)


def test_xlsx_khong_bi_excel_an_kieu() -> None:
    """Bản .xlsx phải là ẢNH CHỤP TRUNG THỰC của labels.csv, không để Excel suy kiểu.

    Kho này đã nhiều lần chảy máu vì ép kiểu: label_in_train '1' -> 1.0, âm Quốc ngữ THẬT
    "nan" -> NaN, crop_w '138' -> 138.0. Excel suy kiểu còn hăng hơn pandas, nên mặc định
    ở đây là CHUỖI; chỉ vài cột đo được mới ghi kiểu số.
    """
    import csv as _csv
    from pipeline.tools import make_xlsx as mx
    print("[xlsx trung thực với csv]")
    ds = _ds_dir()
    src, dst = ds / "labels.csv", ds / "labels.xlsx"
    if not src.exists() or not dst.exists():
        print("  [bỏ qua] chưa có bộ giao nộp / chưa dựng xlsx"); return
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  [bỏ qua] thiếu openpyxl"); return
    rows = list(_csv.DictReader(open(src, encoding="utf-8")))
    ws = load_workbook(dst, read_only=True)["labels"]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    xl = list(it)
    check(f"số dòng khớp ({len(rows):,})", len(xl) == len(rows), f"xlsx {len(xl):,}")
    check("cột khớp đúng thứ tự", hdr == list(rows[0].keys()))
    i = {c: k for k, c in enumerate(hdr)}
    lech = 0
    for a, b in zip(rows, xl):
        for c in hdr:
            if c in mx.COT_SO:
                continue
            # ô rỗng trong xlsx đọc ra None — phải quy về "" TRƯỚC khi so, nếu không
            # chính phép kiểm lại báo động giả trên mọi ô trống (đã dính: 194.875 ô).
            v = b[i[c]]
            if ("" if v is None else str(v)) != a[c]:
                lech += 1
    check("0 ô lệch trên mọi cột CHUỖI", lech == 0, f"{lech} ô")
    # cột cờ/số-dạng-chuỗi của từng thế hệ: ≤25/08 có label_in_train/crop_w trong labels.csv;
    # 12 cột (16/09) thì image_md5 ('0000e5…' dễ mất số 0 đầu) và bbox ('[…]') phải nguyên.
    if "label_in_train" in i:
        lit = {str(r[i["label_in_train"]]) for r in xl if r[i["label_in_train"]] is not None}
        check("label_in_train KHÔNG bị hoá số ('1.0')", lit <= {"0", "1"}, str(sorted(lit))[:50])
    if "crop_w" in i:
        cw = {str(r[i["crop_w"]]) for r in xl if r[i["crop_w"]]}
        check("crop_w không có đuôi '.0'", not any("." in v for v in cw), str(sorted(cw)[:3]))
    if "image_md5" in i:
        md = [str(r[i["image_md5"]]) for r in xl if r[i["image_md5"]] is not None]
        check("image_md5 giữ đủ 12 hex (không mất số 0 đầu / không hoá số)",
              all(len(v) == 12 for v in md), str([v for v in md if len(v) != 12][:3]))
    if "bbox" in i:
        check("bbox giữ nguyên chuỗi '[x1, y1, x2, y2]'",
              all(str(r[i["bbox"]]).startswith("[") for r in xl))
    nan = [r for r in rows if str(r["syllable"]).lower() == "nan"]
    if nan:
        gnan = [r for r in xl if str(r[i["syllable"]]).lower() == "nan"]
        check(f"âm Quốc ngữ 'nan' còn nguyên ({len(nan)} ô)", len(gnan) == len(nan),
              f"xlsx còn {len(gnan)}")


def test_batch_by_rule() -> None:
    """Mẻ chấm PHẢI tách GOLD thành direct/similar — hai lớp rủi ro khác nhau."""
    from pathlib import Path as _P
    import json as _json, collections as _c
    REPO = _P(__file__).resolve().parents[2]
    print("[mẻ chấm — tách theo LUẬT]")
    src = (REPO / "pipeline" / "ground_truth" / "make_combined_batch.py").read_text(encoding="utf-8")
    check("có cờ --by-rule", "--by-rule" in src)
    check("tách được s1_inter_s2_similar", "s1_inter_s2_similar" in src)
    m = REPO / "dataset_out" / "human_audit" / "audit_combined" / "manifest.jsonl"
    if not m.exists():
        print("  [bỏ qua] chưa dựng mẻ"); return
    rows = [_json.loads(l) for l in open(m, encoding="utf-8")]
    st = {str(r.get("stratum", "")).split("|")[0] for r in rows}
    for want in ("GOLD/direct", "GOLD/similar", "SYLLABLE"):
        check(f"có tầng {want}", want in st)
    # BẤT BIẾN ĐÚNG: ô MẪU phải có trọng số; ô LẶP ẨN phải KHÔNG có.
    # Ô lặp dùng đo κ nội tại (người chấm có tự nhất quán không), KHÔNG phải quan sát
    # độc lập — gán trọng số cho nó là ĐẾM TRÙNG dân số và thổi ước lượng.
    mau = [r for r in rows if not r.get("repeat_of")]
    lap = [r for r in rows if r.get("repeat_of")]
    check(f"{len(mau)} ô MẪU đều có design_weight (Horvitz-Thompson)",
          all(r.get("design_weight") for r in mau))
    check(f"{len(lap)} ô LẶP ẨN đều KHÔNG có design_weight (tránh đếm trùng)",
          all(not r.get("design_weight") for r in lap))
    # cảnh báo ảnh hỏng phải tới được manifest
    check("manifest mang crop_quality_flag", "crop_quality_flag" in rows[0])
    h = (m.parent / "audit.html").read_text(encoding="utf-8")
    nbad = sum(1 for r in rows if r.get("crop_quality_flag") in ("blank", "truncated"))
    if nbad:
        check(f"{nbad} ô ảnh hỏng CÓ cảnh báo hiện ra trong HTML",
              "ẢNH TRẮNG" in h or "ẢNH BỊ CẮT" in h)


def test_doi_soat_the_he_v3() -> None:
    """A-13 (N15 B1–B8): khối v3 của doi_soat_the_he trên dữ liệu giả có đáp án.

    Dựng 2 cột × 12 chữ Nôm ở cả hai đời, rồi chèn đúng MỘT ca cho mỗi lớp: đổi ghép
    (syl_idx), ô mất, ô mới, bbox đổi, md5 đổi, GOLD→REVIEW, L1 đổi/thua/hoà, khe giả,
    lệch chéo, QĐ-01 (locked + hàng xóm đổi). Mỗi số đo phải bằng đúng số ca đã chèn.
    """
    print("[tools.doi_soat_the_he v3]")
    import json, tempfile
    from pipeline.tools import doi_soat_the_he as D

    def bbox(i):
        return json.dumps([100, 100 + 110 * i, 200, 200 + 110 * i])

    def cot(col, n, tier="GOLD", rule="s1_inter_s2_direct"):
        rows = []
        for i in range(n):
            rows.append({"book": "b", "page": "p", "column": str(col), "nom_idx": str(i), "syl_idx": str(i),
                         "tier": tier, "rule": rule, "label": chr(0x4E00 + i), "syllable": f"a{i}",
                         "bbox": bbox(i), "image_md5": f"m{col}{i:02d}", "ocr_char": chr(0x4E00 + i)})
        return rows
    cu = pd.DataFrame(cot(1, 12) + cot(2, 12))
    moi = pd.DataFrame(cot(1, 12) + cot(2, 12))
    for c, v in (("count_source", "equal_qn"), ("box_source", "detector"), ("l1_support", "0"),
                 ("l1_tie", "0"), ("tier_v3", "CHAR_A"), ("n_ocr", "12"), ("n_qn", "12"),
                 ("syllable_ocr", ""), ("qd01_locked", "0"), ("qd01_excluded", "0")):
        moi[c] = v
    moi["syllable_ocr"] = moi["syllable"]
    m = moi.set_index(["column", "nom_idx"])
    # B2: cột 1 nom 3 ghép sang âm khác (syl_idx 3 -> 4)
    m.loc[("1", "3"), "syl_idx"] = "4"
    # B3: cột 1 nom 11 MẤT (xoá), cột 2 thêm nom 12 MỚI
    # B4/B5: cột 2 nom 2 bbox + md5 đổi (usable)
    m.loc[("2", "2"), "bbox"] = json.dumps([101, 320, 201, 420]); m.loc[("2", "2"), "image_md5"] = "zz"
    # B7: cột 2 nom 5 GOLD -> REVIEW
    m.loc[("2", "5"), "tier"] = "REVIEW"; m.loc[("2", "5"), "rule"] = "no_context"; m.loc[("2", "5"), "label"] = ""
    # B8: cột 2 nom 6 L1 đổi (support 2), nom 7 thua (-1), nom 8 hoà (tie)
    m.loc[("2", "6"), "rule"] = D.RULE_L1; m.loc[("2", "6"), "l1_support"] = "2"; m.loc[("2", "6"), "syllable"] = "á6"
    m.loc[("2", "7"), "l1_support"] = "-1"
    m.loc[("2", "8"), "l1_tie"] = "1"
    # khe giả: cột 1 mất nom 11 -> 11 ô ghép < n_ocr 12 (n_ocr == n_qn == 12 >= 10); lệch chéo = nom 3
    moi = m.reset_index()
    moi = moi[~((moi["column"] == "1") & (moi["nom_idx"] == "11"))]
    moi = pd.concat([moi, pd.DataFrame([{**moi.iloc[-1].to_dict(), "column": "2", "nom_idx": "12", "syl_idx": "12",
                                          "bbox": bbox(12), "image_md5": "new", "tier": "SYLLABLE",
                                          "rule": "syl_ctx:bigram", "label": ""}])], ignore_index=True)
    # QĐ-01: cột 2 nom 9 khoá; hàng xóm prev đúng, next (nom 10) trong bản mới đổi bbox
    moi.loc[(moi["column"] == "2") & (moi["nom_idx"] == "9"), ["rule", "label", "qd01_locked"]] = \
        ["quyet_dinh_nguoi:qd01_cell_lock", "𠊚", "1"]
    moi.loc[(moi["column"] == "2") & (moi["nom_idx"] == "10"), "bbox"] = json.dumps([100, 1205, 200, 1305])
    qd = pd.DataFrame([{"book": "b", "page": "p", "column": "2", "nom_idx": "9", "label": "𠊚",
                        "bbox_cu": bbox(9), "prev_bbox_cu": bbox(8), "next_bbox_cu": bbox(10),
                        "image_md5_cu": "m209"}])
    with tempfile.TemporaryDirectory() as td:
        L, R = D.doi_soat_v3(cu, moi, qd, Path(td))
        check("B1 ghép/mất/mới = 23/1/1", (R["n_ghep"], R["n_mat"], R["n_moi"]) == (23, 1, 1), str((R["n_ghep"], R["n_mat"], R["n_moi"])))
        check("B2 đổi ghép đúng 1 ô GOLD", R["b2_doi_ghep"] == {"GOLD": 1}, str(R["b2_doi_ghep"]))
        check("B3 GOLD mất 1", R["b3_gold_mat"] == 1)
        check("B4 bbox đổi 2 (1 usable + hàng xóm QĐ-01)", R["b4_bbox_doi"] == 2 and R["b4_bbox_doi_usable"] == 2, str(R["b4_bbox_doi"]))
        check("B5 md5 đổi 1 usable, QĐ-01 0", R["b5_md5_doi_usable"] == 1 and R["b5_md5_doi_qd01"] == 0)
        check("B5a hàng xóm: prev khớp, next khác", R["b5a"]["prev_khac"] == 0 and R["b5a"]["next_khac"] == 1, str(R["b5a"]))
        check("B7 GOLD→REVIEW 1, rule direct", R["b7_tong"] == 1 and R["b7_rule_cu"] == {"s1_inter_s2_direct": 1})
        check("B8 đổi/thua/hoà = 1/1/1, ghi đè 0",
              R["b8"].get("đổi (l1_support > 0)") == 1 and R["b8"].get("giữ gốc (thua)") == 1
              and R["b8"].get("hoà (giữ gốc)") == 1 and R["b8_ghi_de"] == 0, str(R["b8"]))
        check("khe giả 1/2 cột, lệch chéo 1", R["khe_gia_cot"] == 1 and R["khe_gia_tren"] == 2 and R["lech_cheo"] == 1,
              str((R["khe_gia_cot"], R["khe_gia_tren"], R["lech_cheo"])))
        check("usable cũ 24 -> mới 23", R["usable_cu"] == 24 and R["usable_moi"] == 23)
        check("CSV chi tiết được ghi", all((Path(td) / f).exists() for f in
                                           ("B2_doi_am_ghep.csv", "B3_o_mat.csv", "B5_md5_doi.csv", "B7_gold_sang_review.csv")))
        bkv = D.bang_ky_vong(R, None)
        check("bảng kỳ vọng có dòng usable + B8 ghi đè", any("Ô dùng được" in l for l in bkv) and any("ghi đè bản in" in l for l in bkv))


def main() -> int:
    print("=" * 64)
    print("TOOLS SELFTEST")
    print("=" * 64)
    test_fix_tone()
    test_sem_score()
    test_syllable_normalize()
    test_qn_charfix()
    test_check_deps()
    test_dict_candidates()
    test_variant_table()
    test_dataset_docs()
    test_proto_index_ro_ri_split()
    test_publish_doc_csv_phong_ve()
    test_step1_khong_mat_trang()
    test_co_trang_lech_cot()
    test_run_pipeline_grep_dem()
    test_xlsx_khong_bi_excel_an_kieu()
    test_batch_by_rule()
    test_doi_soat_the_he_v3()
    print("=" * 64)
    print(f"RESULT: {_passed} passed, {_failed} failed")
    print("=" * 64)
    return 1 if _failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
